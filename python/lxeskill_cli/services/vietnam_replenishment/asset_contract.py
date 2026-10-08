"""SKU map input contracts and the generated Vietnam workbook layout."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
from typing import Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .numeric_contract import WorkbookInputError, excel_number


REQUIRED_SHEETS = (
    "越南备货清单",
    "雅仓库存",
    "雅仓动销",
    "数据更改",
    "库存商品信息",
)

MAIN_HEADERS = {
    "B1": "判断热销",
    "E1": "SKU",
    "G1": "成本",
    "H1": "陆运",
    "AA1": "上架时间",
    "AE1": "跨境价",
    "AJ1": "折扣价",
    "AN1": "总在途",
    "AV1": "30天",
    "AW1": "15天",
    "AX1": "7天",
    "AY1": "汇率",
}

AUXILIARY_HEADERS = {
    "雅仓库存": {"B1": "SKU", "F1": "库存数量", "H1": "在途数量", "J1": "可用库存"},
    "雅仓动销": {"A1": "SKU", "E1": "7天销量", "F1": "15天销量", "G1": "30天销量"},
    "数据更改": {"A1": "30天", "B1": "15天", "C1": "7天", "D1": "汇率"},
    "库存商品信息": {"B1": "SKU", "K1": "创建时间"},
}


class AssetContractError(ValueError):
    """A supplied workbook does not meet the SKU map input contract."""


@dataclass(frozen=True)
class SkuMapIssue:
    sheet: str
    row: int
    column: int | None
    sku: str | None
    field: str
    reason: str

    def describe(self) -> str:
        column = get_column_letter(self.column) if self.column else ""
        sku = f" · SKU {self.sku}" if self.sku is not None else ""
        return f"{self.sheet}!{column}{self.row}{sku} · {self.field}: {self.reason}"


class SkuMapValidationError(AssetContractError):
    """All independently detectable issues in the header or data rows."""

    def __init__(self, path: str | Path, issues: list[SkuMapIssue], *, headers: bool = False):
        self.path = Path(path)
        self.issues = tuple(sorted(issues, key=lambda issue: (issue.row, issue.column or 0)))
        self.headers = headers
        rows = len({issue.row for issue in issues})
        self.summary = (
            f"SKU 映射表表头发现 {len(issues)} 个错误；请先修正表头，本次未检查数据行。"
            if headers else f"SKU 映射表共发现 {len(issues)} 个错误，涉及 {rows} 行；整表校验未通过。"
        )
        super().__init__(self.summary + "\n" + "\n".join(issue.describe() for issue in self.issues))


@dataclass(frozen=True)
class SkuParameters:
    cost: Decimal | None = None
    cross_border_price: Decimal | None = None
    discount_price: Decimal | None = None
    hot_flag: int | None = None
    listed_at: str | None = None


_PARAMETER_HEADERS = ("SKU", "成本", "跨境价", "折扣价", "热销标记")


def _blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def load_sku_parameters(path: str | Path) -> dict[str, SkuParameters]:
    """Parse first-sheet inputs without filling gaps.

    Runtime callers must use validate_usable_sku_parameters for completeness.
    """
    return _load_sku_parameters(path, complete=False)


def _load_sku_parameters(path: str | Path, *, complete: bool) -> dict[str, SkuParameters]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise AssetContractError(f"读取 SKU 参数表失败: {type(exc).__name__}: {exc}") from exc

    try:
        sheet = workbook.worksheets[0]
        name = sheet.title
        first_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
        columns: dict[str, int] = {}
        issues: list[SkuMapIssue] = []
        for index, cell in enumerate(first_row, 1):
            if not isinstance(cell.value, str) or not cell.value.strip():
                continue
            header = cell.value.strip()
            if header in columns:
                issues.append(SkuMapIssue(name, 1, index, None, header, f"重复表头 {header!r}"))
            else:
                columns[header] = index
        for header in _PARAMETER_HEADERS:
            if header not in columns:
                issues.append(SkuMapIssue(name, 1, None, None, header, f"缺少表头 {header!r}"))
        if issues:
            raise SkuMapValidationError(path, issues, headers=True)

        def cell_at(row: tuple, row_number: int, header: str):
            column = columns[header]
            cell = row[column - 1]
            coordinate = f"{name}!{get_column_letter(column)}{row_number}"
            if cell.data_type == "f":
                raise AssetContractError(f"输入格不能是公式，实际为 {cell.value!r}")
            return cell.value, coordinate

        def number_at(row: tuple, row_number: int, header: str, sku: str | None) -> Decimal | None:
            value, _ = cell_at(row, row_number, header)
            if complete and header != "热销标记":
                try:
                    return excel_number(value, sku or "（无有效 SKU）", header, positive=True)
                except ArithmeticError as exc:
                    raise AssetContractError(f"数值无法用于 Excel: {type(exc).__name__}: {exc}") from exc
            if _blank(value):
                if complete:
                    raise AssetContractError("热销标记必填，且只能是 1 或 2")
                return None
            if isinstance(value, bool):
                raise AssetContractError("必须是有限非负数，实际为布尔值")
            try:
                result = Decimal(str(value))
            except (InvalidOperation, ValueError, TypeError) as exc:
                raise AssetContractError(
                    f"不是数字: {type(exc).__name__}: {exc}"
                ) from exc
            if not result.is_finite() or result < 0:
                raise AssetContractError("必须是有限非负数")
            if header == "热销标记" and result not in (1, 2):
                raise AssetContractError("热销标记只能是 1 或 2")
            return result

        def listed_at_value(row: tuple, row_number: int) -> str | None:
            listed_at = None
            if "上架时间" in columns:
                listed_value, _ = cell_at(row, row_number, "上架时间")
                if not _blank(listed_value):
                    if isinstance(listed_value, str):
                        listed_at = listed_value.strip()
                        if not re.match(r"^\d{4}-\d{2}-\d{2}(?:$|[ T])", listed_at):
                            raise AssetContractError(
                                "上架时间文本必须以 ISO YYYY-MM-DD 开头"
                            )
                        try:
                            if len(listed_at) == 10:
                                date.fromisoformat(listed_at)
                            else:
                                datetime.fromisoformat(listed_at)
                        except ValueError as exc:
                            raise AssetContractError(
                                f"上架时间日期或时间无效: {type(exc).__name__}: {exc}"
                            ) from exc
                    elif isinstance(listed_value, date):
                        listed_cell = row[columns["上架时间"] - 1]
                        listed_at = (
                            listed_value.date().isoformat()
                            if isinstance(listed_value, datetime)
                            and listed_cell.number_format == "yyyy-mm-dd"
                            else str(listed_value).strip()
                        )
                    else:
                        raise AssetContractError("上架时间必须是文本或日期")
            return listed_at

        result: dict[str, SkuParameters] = {}
        first_rows: dict[str, int] = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(_blank(cell.value) for cell in row):
                continue
            row_issue_count = len(issues)
            sku = None
            try:
                sku_value, _ = cell_at(row, row_number, "SKU")
                if not isinstance(sku_value, str) or not sku_value.strip():
                    raise AssetContractError("SKU 必须是非空文本")
                sku = sku_value.strip()
                if sku in first_rows:
                    original = f"{name}!{get_column_letter(columns['SKU'])}{first_rows[sku]}"
                    raise AssetContractError(f"SKU 与 {original} 重复")
                # Even an otherwise invalid row reserves its SKU for duplicate detection.
                first_rows[sku] = row_number
            except AssetContractError as exc:
                issues.append(SkuMapIssue(name, row_number, columns["SKU"], sku, "SKU", str(exc)))

            values = {}
            for field, header in (
                ("cost", "成本"), ("cross_border_price", "跨境价"),
                ("discount_price", "折扣价"), ("hot_flag", "热销标记"),
            ):
                try:
                    value = number_at(row, row_number, header, sku)
                    values[field] = int(value) if field == "hot_flag" and value is not None else value
                except (AssetContractError, WorkbookInputError) as exc:
                    issues.append(SkuMapIssue(name, row_number, columns[header], sku, header, str(exc)))
            try:
                values["listed_at"] = listed_at_value(row, row_number)
            except AssetContractError as exc:
                issues.append(SkuMapIssue(name, row_number, columns["上架时间"], sku, "上架时间", str(exc)))
            if len(issues) == row_issue_count:
                result[sku] = SkuParameters(**values)
        if issues:
            raise SkuMapValidationError(path, issues)
        if complete and not result:
            raise AssetContractError("当前越南 SKU 参数映射表没有 SKU，请重新上传")
        return result
    except AssetContractError:
        raise
    except Exception as exc:
        raise AssetContractError(f"读取 SKU 参数表失败: {type(exc).__name__}: {exc}") from exc
    finally:
        workbook.close()


def validate_sku_parameter_values(values: Mapping[str, SkuParameters]) -> None:
    """Every supplied SKU must be complete; ERP coverage is not required."""
    if not values:
        raise WorkbookInputError("当前越南 SKU 参数映射表没有 SKU，请重新上传")
    for sku, row in values.items():
        if not isinstance(row, SkuParameters):
            raise WorkbookInputError(f"SKU {sku} 的运营 SKU 映射行无效")
        for field, label in (
            ("cost", "成本"),
            ("cross_border_price", "跨境价"),
            ("discount_price", "折扣价"),
        ):
            excel_number(getattr(row, field), sku, label, positive=True)
        if isinstance(row.hot_flag, bool) or row.hot_flag not in (1, 2):
            raise WorkbookInputError(f"SKU {sku} 的热销标记必填，且只能是 1 或 2")


def validate_usable_sku_parameters(path: str | Path) -> dict[str, SkuParameters]:
    """Require complete mapped SKUs with positive Excel-exact prices."""
    return _load_sku_parameters(path, complete=True)


__all__ = [
    "AssetContractError",
    "SkuMapIssue",
    "SkuMapValidationError",
    "SkuParameters",
    "load_sku_parameters",
    "validate_usable_sku_parameters",
    "validate_sku_parameter_values",
]
