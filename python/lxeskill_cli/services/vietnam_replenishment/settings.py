"""App-wide Vietnam inputs, shared by Desktop and the ordinary lxeskill command.

The module entrypoint is a private Desktop bridge, not a model command. All
writes are atomic and serialized; the calculation keeps its own input snapshot.
"""
from __future__ import annotations

from contextlib import contextmanager
from dataclasses import asdict
from datetime import datetime, timezone
from decimal import Decimal
import json
import os
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from uuid import uuid4

from shared.process_lock import interprocess_lock
from shared.repository import state_root
from shared.workspace import resolve_workspace_input
from services.yacang.errors import safe_remote_detail
from .sku_map_store import SkuMapStoreError, validate_sku_map
from .validation_diagnostics import validation_diagnostic
from .workbook import RecommendationConfig, validate_recommendation_config

_FIELDS = tuple(asdict(RecommendationConfig()))
_DECIMAL = re.compile(r"^[+]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?$")


def data_directory() -> Path:
    return state_root() / "skill-data" / "vietnam-stock-recommendation"


def config_json(config: RecommendationConfig) -> dict[str, str]:
    return {key: str(value) for key, value in asdict(config).items()}


def parse_parameters(raw: object) -> RecommendationConfig:
    if not isinstance(raw, dict) or set(raw) != set(_FIELDS):
        raise ValueError(f"越南备货参数必须包含且仅包含: {', '.join(_FIELDS)}")
    values = {}
    for key in _FIELDS:
        value = raw[key]
        if not isinstance(value, str) or not 0 < len(value.strip()) <= 128 or not _DECIMAL.fullmatch(value.strip()):
            raise ValueError(f"{key} 必须是非负十进制数字文本")
        values[key] = Decimal(value.strip())
    result = RecommendationConfig(**values)
    validate_recommendation_config(result)
    return result


def _json(path: Path) -> object:
    if path.stat().st_size > 1024 * 1024:
        raise ValueError(f"配置文件超过 1 MiB: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def _write_parameters(root: Path, config: RecommendationConfig) -> None:
    temporary = root / f".parameters-{uuid4().hex}.tmp"
    try:
        with temporary.open("x", encoding="utf-8") as stream:
            json.dump(config_json(config), stream, ensure_ascii=False, indent=2)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, root / "parameters.json")
    finally:
        temporary.unlink(missing_ok=True)


def _initialize_parameters(root: Path) -> RecommendationConfig:
    path = root / "parameters.json"
    if path.exists():
        return parse_parameters(_json(path))
    config = RecommendationConfig()
    _write_parameters(root, config)
    return config


def _copy_validated(source: Path, destination: Path) -> None:
    expected = validate_sku_map(source, label="SKU 映射表")
    with source.open("rb") as reader, destination.open("xb") as writer:
        count = 0
        while block := reader.read(1024 * 1024):
            count += len(block)
            if count > 20 * 1024 * 1024:
                raise SkuMapStoreError("复制中的 SKU 映射表超过 20 MiB")
            writer.write(block)
        writer.flush()
        os.fsync(writer.fileno())
    if validate_sku_map(destination, label="SKU 映射表快照") != expected:
        raise SkuMapStoreError("SKU 映射表在复制期间发生变化，请重新选择")


def _replace_map(root: Path, source: Path) -> None:
    temporary = root / f".sku-map-{uuid4().hex}.xlsx"
    try:
        _copy_validated(source, temporary)
        os.replace(temporary, root / "sku-map.xlsx")
    finally:
        temporary.unlink(missing_ok=True)


@contextmanager
def _locked():
    root = data_directory()
    root.mkdir(parents=True, exist_ok=True)
    with interprocess_lock(root / ".settings.lock"):
        yield root


def read_state() -> dict:
    result = {"directory": str(data_directory()), "parameters": None, "parameters_error": None,
              "sku_map": None, "sku_map_error": None}
    with _locked() as root:
        try:
            result["parameters"] = config_json(_initialize_parameters(root))
        except Exception as exc:
            result["parameters_error"] = safe_remote_detail(f"{type(exc).__name__}: {exc}")
        try:
            target = root / "sku-map.xlsx"
            if target.exists():
                size, _ = validate_sku_map(target, label="已保存的 SKU 映射表")
                result["sku_map"] = {"path": str(target), "file_name": target.name, "size_bytes": size,
                    "updated_at": datetime.fromtimestamp(target.stat().st_mtime, timezone.utc).isoformat()}
        except Exception as exc:
            result["sku_map_error"] = validation_diagnostic(exc)
    return result


def save_parameters(raw: object) -> dict:
    config = parse_parameters(raw)
    with _locked() as root:
        _write_parameters(root, config)
    return read_state()


def upload_map(source: Path) -> dict:
    with _locked() as root:
        _replace_map(root, source)
    return read_state()


def _export_destination(destination: Path) -> None:
    if not destination.is_absolute() or destination.suffix.lower() != ".xlsx":
        raise ValueError("导出路径必须是绝对路径且以 .xlsx 结尾")
    root = data_directory().resolve()
    source = root / "sku-map.xlsx"
    if (destination.resolve().is_relative_to(root)
            or destination.exists() and source.exists() and destination.samefile(source)):
        raise ValueError("导出不能覆盖应用内部保存的 SKU 映射表，请选择其他位置")


def _write_map_template(destination: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    try:
        sheet = book.active
        sheet.title = "SKU参数映射"
        sheet.append(["SKU", "热销标记", "成本", "跨境价", "折扣价"])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:E1"
        sheet.row_dimensions[1].height = 26
        for column, width in zip("ABCDE", (28, 16, 18, 18, 18)):
            sheet.column_dimensions[column].width = width
            sheet[f"{column}1"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            sheet[f"{column}1"].fill = PatternFill("solid", fgColor="1F4E78")
        sheet.column_dimensions["A"].number_format = "@"
        sheet["A2"].number_format = "@"
        sheet["A1"].comment = Comment("必填文本，保留前导零。每个已填写的 SKU 都须填写其余四项；完全空白的行会跳过。", "LXE Agent")
        sheet["B1"].comment = Comment("必填，只能填写 1 或 2，不设默认值。", "LXE Agent")
        for column in "CDE":
            sheet[f"{column}1"].comment = Comment("必填真实数值，必须大于 0，不接受空白、零或公式。", "LXE Agent")
        with destination.open("xb") as stream:
            book.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
    finally:
        book.close()


def export_map(kind: str, destination: Path) -> dict:
    """Export without loading parameters or altering the saved map.

    The native save dialog owns overwrite confirmation. Only replace the selected
    destination after the complete template or validated snapshot is ready.
    """
    if kind not in ("template", "current"):
        raise ValueError("无效的 SKU 映射表导出类型")
    _export_destination(destination)
    temporary = destination.parent / f".vietnam-export-{uuid4().hex}.xlsx"
    try:
        if kind == "template":
            _write_map_template(temporary)
        else:
            with _locked() as root:
                _copy_validated(root / "sku-map.xlsx", temporary)
        _export_destination(destination)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)
    return {"path": str(destination)}


@contextmanager
def run_inputs(sku_map: str | None = None):
    """Snapshot settings and map; release the lock before calculation."""
    with TemporaryDirectory(prefix="vietnam-inputs-") as temporary:
        snapshot = Path(temporary) / "sku-map.xlsx"
        with _locked() as root:
            config = _initialize_parameters(root)
            source = resolve_workspace_input(sku_map) if sku_map is not None else root / "sku-map.xlsx"
            if sku_map is None and not source.exists():
                raise SkuMapStoreError("请在越南备货设置上传 SKU 映射表，或用 --sku-map 指定文件")
            _copy_validated(source, snapshot)
        yield snapshot, config, str(root / "parameters.json"), str(source)


def main() -> int:
    import sys
    try:
        action = sys.argv[1]
        if action == "read" and len(sys.argv) == 2:
            result = read_state()
        elif action == "save" and len(sys.argv) == 3:
            result = save_parameters(json.loads(sys.argv[2]))
        elif action == "upload" and len(sys.argv) == 3:
            source = Path(sys.argv[2])
            if not source.is_absolute():
                raise ValueError("上传文件必须是绝对路径")
            result = upload_map(source)
        elif action == "export" and len(sys.argv) == 4:
            result = export_map(sys.argv[2], Path(sys.argv[3]))
        else:
            raise ValueError("无效的桌面设置操作")
        print(json.dumps({"success": True, "data": result}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"success": False, "error": validation_diagnostic(exc)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
