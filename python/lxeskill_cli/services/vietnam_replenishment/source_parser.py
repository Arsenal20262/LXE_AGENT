"""Classify and validate three existing VN8806 XLSX reports locally."""

from __future__ import annotations

from contextlib import closing
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import os
from pathlib import Path
from typing import Mapping

from openpyxl import load_workbook

from services.yacang.errors import safe_remote_detail
from services.yacang.validation import (
    INVENTORY_LIST_HEADERS,
    INVENTORY_SALES_HEADERS,
    WAREHOUSE_PRODUCTS_HEADERS,
    validate_inventory_list_workbook,
    validate_inventory_sales_workbook,
    validate_warehouse_products_workbook,
)
from shared.filesystem import filesystem_path


VIETNAM_REPORTS = (
    "inventory-sales",
    "inventory-current-snapshot",
    "warehouse-products",
)
_HEADERS = {
    "inventory-sales": INVENTORY_SALES_HEADERS,
    "inventory-current-snapshot": INVENTORY_LIST_HEADERS,
    "warehouse-products": WAREHOUSE_PRODUCTS_HEADERS,
}


class VietnamSourceError(ValueError):
    """The three current Yacang sources cannot be used as one VN run."""


@dataclass(frozen=True)
class VietnamSources:
    skus: tuple[str, ...]
    sales: Mapping[str, Mapping[str, object]]
    inventory: Mapping[str, Mapping[str, object]]
    products: Mapping[str, Mapping[str, object]]
    missing_sales: tuple[str, ...]
    missing_inventory: tuple[str, ...]
    missing_products: tuple[str, ...]
    in_transit: Mapping[str, Decimal | None]
    missing_in_transit: tuple[str, ...]
    in_transit_mismatch: tuple[str, ...]
    artifacts: Mapping[str, Path]


def _paths(artifacts: object) -> dict[str, Path]:
    try:
        entries = list(artifacts)
    except TypeError as exc:
        raise VietnamSourceError("雅仓导出必须包含三份文件元数据") from exc
    if len(entries) != 3:
        raise VietnamSourceError(f"雅仓导出必须恰好包含三份文件，实际 {len(entries)} 份")

    paths: dict[str, Path] = {}
    for entry in entries:
        if not isinstance(entry, Mapping):
            raise VietnamSourceError("雅仓导出文件元数据不是对象")
        report = entry.get("report")
        if report not in VIETNAM_REPORTS or report in paths:
            raise VietnamSourceError(f"雅仓导出报表缺失、重复或不受支持: {report!r}")
        expected_warehouse = None if report == "warehouse-products" else "VN8806"
        if entry.get("warehouse") != expected_warehouse:
            raise VietnamSourceError(
                f"{report} 仓库必须为 {expected_warehouse or '全局'}，实际为 {entry.get('warehouse')!r}; "
                "本轮仅接受 VN8806 两份报表及全局产品资料"
            )
        if entry.get("created_date") is not None:
            raise VietnamSourceError(f"{report} 不得带创建日期筛选，否则会遗漏本轮 SKU")
        raw_path = entry.get("path")
        if not isinstance(raw_path, (str, Path)) or not str(raw_path):
            raise VietnamSourceError(f"{report} 缺少有效文件路径")
        paths[report] = Path(raw_path)

    if set(paths) != set(VIETNAM_REPORTS):
        raise VietnamSourceError(f"雅仓导出报表不齐，实际为 {sorted(paths)}")
    return paths


def _read_rows(path: Path, report: str, *, current_skus: set[str] | None = None) -> dict[str, dict[str, object]]:
    secrets = tuple(os.getenv(name, "") for name in ("LXE_YACANG_MOBILE", "LXE_YACANG_PASSWORD"))

    def diagnostic(value: object) -> str:
        return safe_remote_detail(value, secrets=secrets)

    if report == "inventory-sales":
        validate_inventory_sales_workbook(path, warehouse_code="VN8806", diagnostic=diagnostic)
    elif report == "inventory-current-snapshot":
        validate_inventory_list_workbook(path, warehouse_code="VN8806", diagnostic=diagnostic)
    else:
        validate_warehouse_products_workbook(path, diagnostic=diagnostic)

    headers = _HEADERS[report]
    records: dict[str, dict[str, object]] = {}
    first_rows: dict[str, int] = {}
    with filesystem_path(path).open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            with closing(sheet.iter_rows()) as rows:
                next(rows, ())
                sku_index = headers.index("SKU")
                for row_number, cells in enumerate(rows, 2):
                    if all(cell.value is None for cell in cells):
                        continue
                    sku_cell = cells[sku_index]
                    raw_sku = sku_cell.value
                    if sku_cell.data_type == "f" or not isinstance(raw_sku, str) or not raw_sku.strip():
                        raise VietnamSourceError(f"{report} 第 {row_number} 行 SKU 必须是非空文本")
                    sku = raw_sku.strip()
                    if current_skus is not None and sku not in current_skus:
                        continue
                    for column, cell in enumerate(cells[:len(headers)]):
                        if cell.data_type == "f":
                            raise VietnamSourceError(
                                f"{report} 第 {row_number} 行 {headers[column]} 不能是公式"
                            )
                    if sku in records:
                        raise VietnamSourceError(
                            f"{report} 第 {row_number} 行与第 {first_rows[sku]} 行重复 SKU: {sku}"
                        )
                    values = {
                        header: cells[index].value if index < len(cells) else None
                        for index, header in enumerate(headers)
                    }
                    values["SKU"] = sku
                    records[sku] = values
                    first_rows[sku] = row_number
        finally:
            workbook.close()
    return records


def _nonnegative_number(value: object, *, report: str, sku: str, field: str) -> Decimal | None:
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, bool):
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是有限非负数: 布尔值")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是数字: {type(exc).__name__}: {exc}") from exc
    if not result.is_finite() or result < 0:
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是有限非负数")
    return result


def load_vietnam_sources(artifacts: object) -> VietnamSources:
    """Validate three export artifacts and retain only current VN SKU products."""
    paths = _paths(artifacts)
    try:
        sales = _read_rows(paths["inventory-sales"], "inventory-sales")
        inventory = _read_rows(paths["inventory-current-snapshot"], "inventory-current-snapshot")
        current_skus = set(sales) | set(inventory)
        if not current_skus:
            raise VietnamSourceError("VN8806 库存动销和库存列表均无 SKU，无法确定本轮集合")
        products = _read_rows(paths["warehouse-products"], "warehouse-products", current_skus=current_skus)
    except VietnamSourceError:
        raise
    except Exception as exc:
        raise VietnamSourceError(f"读取本轮雅仓文件失败: {type(exc).__name__}: {exc}") from exc

    in_transit: dict[str, Decimal | None] = {}
    mismatches: list[str] = []
    for sku in sorted(current_skus):
        row = inventory.get(sku)
        authoritative = (
            _nonnegative_number(row["在途数量"], report="inventory-current-snapshot", sku=sku, field="在途数量")
            if row is not None else None
        )
        in_transit[sku] = authoritative
        if sku not in sales:
            continue
        sales_value = _nonnegative_number(
            sales[sku]["在途"], report="inventory-sales", sku=sku, field="在途"
        )
        if authoritative is not None and sales_value is not None and sales_value != authoritative:
            mismatches.append(sku)

    return VietnamSources(
        skus=tuple(sorted(current_skus)),
        sales=sales,
        inventory=inventory,
        products=products,
        missing_sales=tuple(sorted(current_skus - set(sales))),
        missing_inventory=tuple(sorted(current_skus - set(inventory))),
        missing_products=tuple(sorted(current_skus - set(products))),
        in_transit=in_transit,
        missing_in_transit=tuple(sku for sku in sorted(current_skus) if in_transit[sku] is None),
        in_transit_mismatch=tuple(sorted(mismatches)),
        artifacts=paths,
    )


def _report_headers(path: Path) -> tuple[str, ...]:
    with filesystem_path(path).open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            if len(workbook.sheetnames) != 1:
                raise VietnamSourceError(f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}")
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            with closing(sheet.iter_rows(values_only=True)) as rows:
                return tuple(str(value or "").strip() for value in next(rows, ()))
        finally:
            workbook.close()


def classify_vietnam_report_paths(raw_paths: object) -> list[dict[str, object]]:
    """Identify all three local reports by exact headers, never by filename."""
    if not isinstance(raw_paths, (list, tuple)) or len(raw_paths) != 3:
        raise VietnamSourceError("离线生成必须恰好提供三份雅仓报表")
    by_headers = {headers: report for report, headers in _HEADERS.items()}
    seen_paths: set[Path] = set()
    seen_files: set[tuple[int, int]] = set()
    seen_reports: set[str] = set()
    artifacts: list[dict[str, object]] = []
    for index, raw in enumerate(raw_paths, 1):
        if not isinstance(raw, (str, Path)) or not str(raw).strip():
            raise VietnamSourceError(f"第 {index} 份来源缺少有效 .xlsx 路径")
        path = Path(raw)
        if not path.is_absolute() or path.suffix.lower() != ".xlsx":
            raise VietnamSourceError(f"第 {index} 份来源必须是 .xlsx 绝对文件路径")
        try:
            path = path.resolve(strict=True)
            identity = path.stat()
        except OSError as exc:
            raise VietnamSourceError(f"第 {index} 份来源不存在或不可读: {type(exc).__name__}: {exc}") from exc
        if not path.is_file():
            raise VietnamSourceError(f"第 {index} 份来源不是文件")
        file_id = (identity.st_dev, identity.st_ino)
        if path in seen_paths or file_id in seen_files:
            raise VietnamSourceError(f"第 {index} 份来源与另一份文件重复")
        seen_paths.add(path)
        seen_files.add(file_id)
        try:
            headers = _report_headers(path)
        except Exception as exc:
            raise VietnamSourceError(
                f"第 {index} 份 XLSX 无法读取或结构无效: {type(exc).__name__}: {exc}"
            ) from exc
        report = by_headers.get(headers)
        if report is None:
            raise VietnamSourceError(f"第 {index} 份 XLSX 的表头不属于所需三类雅仓报表")
        if report in seen_reports:
            raise VietnamSourceError(f"第 {index} 份 XLSX 的报表类型重复: {report}")
        seen_reports.add(report)
        artifacts.append({
            "report": report,
            "warehouse": None if report == "warehouse-products" else "VN8806",
            "created_date": None,
            "path": str(path),
        })
    if seen_reports != set(VIETNAM_REPORTS):
        raise VietnamSourceError("三份来源必须分别为库存动销、当前库存和仓库产品资料")
    return artifacts


__all__ = [
    "VIETNAM_REPORTS", "VietnamSourceError", "VietnamSources",
    "classify_vietnam_report_paths", "load_vietnam_sources",
]
