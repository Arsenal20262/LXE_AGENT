"""Read supplied VN8806 Yacang reports without fetching or changing source files."""

from __future__ import annotations

from contextlib import closing
from dataclasses import dataclass, field, replace
from decimal import Decimal, InvalidOperation
import hashlib
import os
from pathlib import Path
import shutil
from tempfile import TemporaryDirectory
from typing import Mapping

from openpyxl import load_workbook

from services.yacang.errors import safe_remote_detail
from services.yacang.validation import (
    INVENTORY_LIST_HEADERS,
    INVENTORY_SALES_HEADERS,
    WAREHOUSE_PRODUCTS_HEADERS,
    validate_inventory_list_workbook,
    validate_inventory_sales_workbook,
    validate_warehouse_products_workbook,
)
from shared.filesystem import filesystem_path
from shared.workspace import resolve_workspace_input


_REPORTS = (
    "inventory-sales",
    "inventory-current-snapshot",
    "warehouse-products",
)
_HEADERS = {
    "inventory-sales": INVENTORY_SALES_HEADERS,
    "inventory-current-snapshot": INVENTORY_LIST_HEADERS,
    "warehouse-products": WAREHOUSE_PRODUCTS_HEADERS,
}


class VietnamSourceError(ValueError):
    """The three current Yacang sources cannot be used as one VN run."""


@dataclass(frozen=True)
class VietnamSources:
    skus: tuple[str, ...]
    sales: Mapping[str, Mapping[str, object]]
    inventory: Mapping[str, Mapping[str, object]]
    products: Mapping[str, Mapping[str, object]]
    missing_sales: tuple[str, ...]
    missing_inventory: tuple[str, ...]
    missing_products: tuple[str, ...]
    in_transit: Mapping[str, Decimal | None]
    missing_in_transit: tuple[str, ...]
    in_transit_mismatch: tuple[str, ...]
    artifacts: Mapping[str, Path]
    file_hashes: Mapping[str, str] = field(default_factory=dict)


def _paths(artifacts: object) -> dict[str, Path]:
    try:
        entries = list(artifacts)
    except TypeError as exc:
        raise VietnamSourceError("雅仓导出必须包含三份文件元数据") from exc
    if len(entries) != 3:
        raise VietnamSourceError(f"雅仓导出必须恰好包含三份文件，实际 {len(entries)} 份")

    paths: dict[str, Path] = {}
    for entry in entries:
        if not isinstance(entry, Mapping):
            raise VietnamSourceError("雅仓导出文件元数据不是对象")
        report = entry.get("report")
        if report not in _REPORTS or report in paths:
            raise VietnamSourceError(f"雅仓导出报表缺失、重复或不受支持: {report!r}")
        expected_warehouse = None if report == "warehouse-products" else "VN8806"
        if entry.get("warehouse") != expected_warehouse:
            raise VietnamSourceError(
                f"{report} 仓库必须为 {expected_warehouse or '全局'}，实际为 {entry.get('warehouse')!r}; "
                "本轮仅接受 VN8806 两份报表及全局产品资料"
            )
        if entry.get("created_date") is not None:
            raise VietnamSourceError(f"{report} 不得带创建日期筛选，否则会遗漏本轮 SKU")
        raw_path = entry.get("path")
        if not isinstance(raw_path, (str, Path)) or not str(raw_path):
            raise VietnamSourceError(f"{report} 缺少有效文件路径")
        paths[report] = Path(raw_path)

    if set(paths) != set(_REPORTS):
        raise VietnamSourceError(f"雅仓导出报表不齐，实际为 {sorted(paths)}")
    return paths


def _read_rows(path: Path, report: str, *, current_skus: set[str] | None = None) -> dict[str, dict[str, object]]:
    secrets = tuple(os.getenv(name, "") for name in ("LXE_YACANG_MOBILE", "LXE_YACANG_PASSWORD"))

    def diagnostic(value: object) -> str:
        return safe_remote_detail(value, secrets=secrets)

    if report == "inventory-sales":
        validate_inventory_sales_workbook(path, warehouse_code="VN8806", diagnostic=diagnostic)
    elif report == "inventory-current-snapshot":
        validate_inventory_list_workbook(path, warehouse_code="VN8806", diagnostic=diagnostic)
    else:
        validate_warehouse_products_workbook(path, diagnostic=diagnostic)

    headers = _HEADERS[report]
    records: dict[str, dict[str, object]] = {}
    first_rows: dict[str, int] = {}
    with filesystem_path(path).open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            with closing(sheet.iter_rows()) as rows:
                next(rows, ())
                sku_index = headers.index("SKU")
                for row_number, cells in enumerate(rows, 2):
                    if all(cell.value is None for cell in cells):
                        continue
                    sku_cell = cells[sku_index]
                    raw_sku = sku_cell.value
                    if sku_cell.data_type == "f" or not isinstance(raw_sku, str) or not raw_sku.strip():
                        raise VietnamSourceError(f"{report} 第 {row_number} 行 SKU 必须是非空文本")
                    sku = raw_sku.strip()
                    if current_skus is not None and sku not in current_skus:
                        continue
                    for column, cell in enumerate(cells[:len(headers)]):
                        if cell.data_type == "f":
                            raise VietnamSourceError(
                                f"{report} 第 {row_number} 行 {headers[column]} 不能是公式"
                            )
                    if sku in records:
                        raise VietnamSourceError(
                            f"{report} 第 {row_number} 行与第 {first_rows[sku]} 行重复 SKU: {sku}"
                        )
                    values = {
                        header: cells[index].value if index < len(cells) else None
                        for index, header in enumerate(headers)
                    }
                    values["SKU"] = sku
                    records[sku] = values
                    first_rows[sku] = row_number
        finally:
            workbook.close()
    return records


def _nonnegative_number(value: object, *, report: str, sku: str, field: str) -> Decimal | None:
    if value is None or isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, bool):
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是有限非负数: 布尔值")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as exc:
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是数字: {type(exc).__name__}: {exc}") from exc
    if not result.is_finite() or result < 0:
        raise VietnamSourceError(f"{report} SKU {sku} 的 {field} 不是有限非负数")
    return result


def load_vietnam_sources(artifacts: object) -> VietnamSources:
    """Validate three export artifacts and retain only current VN SKU products."""
    paths = _paths(artifacts)
    try:
        sales = _read_rows(paths["inventory-sales"], "inventory-sales")
        inventory = _read_rows(paths["inventory-current-snapshot"], "inventory-current-snapshot")
        current_skus = set(sales) | set(inventory)
        if not current_skus:
            raise VietnamSourceError("VN8806 库存动销和库存列表均无 SKU，无法确定本轮集合")
        products = _read_rows(paths["warehouse-products"], "warehouse-products", current_skus=current_skus)
    except VietnamSourceError:
        raise
    except Exception as exc:
        raise VietnamSourceError(f"读取本轮雅仓文件失败: {type(exc).__name__}: {exc}") from exc

    in_transit: dict[str, Decimal | None] = {}
    mismatches: list[str] = []
    for sku in sorted(current_skus):
        row = inventory.get(sku)
        authoritative = (
            _nonnegative_number(row["在途数量"], report="inventory-current-snapshot", sku=sku, field="在途数量")
            if row is not None else None
        )
        in_transit[sku] = authoritative
        if sku not in sales:
            continue
        sales_value = _nonnegative_number(
            sales[sku]["在途"], report="inventory-sales", sku=sku, field="在途"
        )
        if authoritative is not None and sales_value is not None and sales_value != authoritative:
            mismatches.append(sku)

    return VietnamSources(
        skus=tuple(sorted(current_skus)),
        sales=sales,
        inventory=inventory,
        products=products,
        missing_sales=tuple(sorted(current_skus - set(sales))),
        missing_inventory=tuple(sorted(current_skus - set(inventory))),
        missing_products=tuple(sorted(current_skus - set(products))),
        in_transit=in_transit,
        missing_in_transit=tuple(sku for sku in sorted(current_skus) if in_transit[sku] is None),
        in_transit_mismatch=tuple(sorted(mismatches)),
        artifacts=paths,
    )


def load_vietnam_files(*, sales: str, inventory: str, products: str) -> VietnamSources:
    """Snapshot and validate local reports; paths do not attest export freshness.

    Warehouse and report identity are checked against file contents. The caller
    must choose complete reports without date filters; file paths alone cannot
    prove that scope or that all three reports came from one export run.
    """
    paths = dict(zip(_REPORTS, (resolve_workspace_input(path) for path in (sales, inventory, products))))
    with TemporaryDirectory(prefix="vietnam-sources-") as temporary:
        artifacts = []
        hashes = {}
        for report, path in paths.items():
            try:
                if not path.is_file() or path.suffix.lower() != ".xlsx":
                    raise VietnamSourceError(f"{report} 必须是存在的 XLSX 文件: {path}")
                snapshot = Path(temporary) / f"{report}.xlsx"
                shutil.copyfile(filesystem_path(path), snapshot)
                with snapshot.open("rb") as stream:
                    hashes[report] = hashlib.file_digest(stream, "sha256").hexdigest()
                artifacts.append({
                    "report": report, "path": snapshot,
                    "warehouse": None if report == "warehouse-products" else "VN8806",
                })
            except (OSError, ValueError) as exc:
                raise VietnamSourceError(f"读取 {report} 文件 {path} 失败: {type(exc).__name__}: {exc}") from exc
        try:
            sources = load_vietnam_sources(artifacts)
        except Exception as exc:
            detail = str(exc)
            for entry in artifacts:
                detail = detail.replace(str(entry["path"]), str(paths[entry["report"]]))
            raise VietnamSourceError(f"校验输入报表失败: {type(exc).__name__}: {detail}") from exc
        return replace(sources, artifacts=paths, file_hashes=hashes)


__all__ = ["VietnamSourceError", "VietnamSources", "load_vietnam_sources", "load_vietnam_files"]
