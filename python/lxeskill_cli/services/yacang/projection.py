from __future__ import annotations

import os
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from services.yacang.errors import YacangError, safe_remote_detail
from services.yacang.validation import validate_inventory_sales_workbook


SALES_MONTHLY_HEADERS = (
    "SKU",
    "商品名",
    "仓库",
    "7天销量",
    "15天销量",
    "30天销量",
)
SALES_90D_HEADERS = (
    "SKU",
    "商品名",
    "仓库",
    "90天销量",
)


def validate_projected_workbook(
    path: Path,
    *,
    expected_headers: tuple[str, ...],
    warehouse_code: str,
) -> int:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - preserve the real parser failure
        raise YacangError("校验业务 XLSX", f"{type(exc).__name__}: {safe_remote_detail(exc)}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise YacangError("校验业务 XLSX", f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        actual_headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if actual_headers != expected_headers:
            raise YacangError("校验业务 XLSX", f"表头不匹配: {safe_remote_detail(actual_headers)}")
        warehouse_index = expected_headers.index("仓库")
        row_count = 0
        wrong_warehouses: set[str] = set()
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            actual = str(row[warehouse_index] or "").strip() if len(row) > warehouse_index else ""
            if actual != warehouse_code and len(wrong_warehouses) < 5:
                wrong_warehouses.add(actual or "[empty]")
        if wrong_warehouses:
            raise YacangError(
                "校验业务 XLSX",
                f"期望仓库 {warehouse_code}，文件包含其他仓库: {sorted(wrong_warehouses)}",
            )
        return row_count
    finally:
        workbook.close()


def project_inventory_sales_workbook(
    source: Path,
    destination: Path,
    *,
    headers: tuple[str, ...],
    warehouse_code: str,
) -> int:
    validate_inventory_sales_workbook(source, warehouse_code=warehouse_code)
    if len(headers) != len(set(headers)) or "仓库" not in headers:
        raise ValueError("业务 XLSX 表头必须唯一且包含仓库")

    source_workbook = load_workbook(source, read_only=True, data_only=True)
    temporary = destination.with_name(f".{destination.stem}.{uuid4().hex}.part.xlsx")
    output_workbook = Workbook(write_only=True)
    try:
        source_sheet = source_workbook[source_workbook.sheetnames[0]]
        rows = source_sheet.iter_rows(values_only=True)
        source_headers = tuple(str(value or "").strip() for value in next(rows, ()))
        missing = [header for header in headers if header not in source_headers]
        if missing:
            raise YacangError("生成业务 XLSX", f"原始文件缺少字段: {safe_remote_detail(missing)}")
        indexes = [source_headers.index(header) for header in headers]

        output_sheet = output_workbook.create_sheet("数据")
        output_sheet.append(headers)
        for row in rows:
            if any(value is not None for value in row):
                output_sheet.append(tuple(row[index] if index < len(row) else None for index in indexes))

        destination.parent.mkdir(parents=True, exist_ok=True)
        output_workbook.save(temporary)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        source_workbook.close()
        output_workbook.close()

    try:
        row_count = validate_projected_workbook(
            temporary,
            expected_headers=headers,
            warehouse_code=warehouse_code,
        )
        os.replace(temporary, destination)
        return row_count
    finally:
        temporary.unlink(missing_ok=True)


__all__ = [
    "SALES_90D_HEADERS",
    "SALES_MONTHLY_HEADERS",
    "project_inventory_sales_workbook",
    "validate_projected_workbook",
]
