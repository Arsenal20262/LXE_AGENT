from __future__ import annotations

import json
import os
import time
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook

from services.yacang.client import YacangClient
from services.yacang.errors import YacangError, safe_remote_detail
from shared.datasets import dataset_dir


WAREHOUSES: tuple[tuple[str, int], ...] = (
    ("MY8801", 26),
    ("PH8805", 46),
    ("TH8802", 47),
    ("VN8806", 80),
)
EXPECTED_HEADERS = (
    "SKU", "商品名", "仓库", "3天销量", "7天销量", "15天销量", "30天销量", "60天销量",
    "90天销量", "库存", "占用", "在途", "冻结", "可用", "缺货数量", "创建日期",
)
EXPORT_NAME = "库存动销导出"
EXPORT_TYPE = "10"
DEFAULT_POLL_INTERVAL_SECONDS = 5.0
DEFAULT_TIMEOUT_SECONDS = 180.0
DEFAULT_CACHE_MAX_AGE_SECONDS = 300.0


@dataclass(frozen=True)
class ExportRequest:
    warehouse_code: str
    warehouse_id: int
    start_date: str
    end_date: str

    @property
    def filename(self) -> str:
        return f"yacang_inventory_sales_{self.warehouse_code}_{self.start_date}_{self.end_date}.xlsx"


def _iso_date(value: Any, name: str) -> str:
    text = str(value or "").strip()
    try:
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:
        raise ValueError(f"{name} 必须是 YYYY-MM-DD，收到: {text or '[empty]'}") from exc


def validate_date_range(start_date: Any, end_date: Any) -> tuple[str, str]:
    start = _iso_date(start_date, "start_date")
    end = _iso_date(end_date, "end_date")
    if start > end:
        raise ValueError(f"start_date 不能晚于 end_date: {start} > {end}")
    return start, end


def _walk_values(value: Any) -> Iterable[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, child in value.items():
            yield str(key), child
            yield from _walk_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_values(child)


def task_matches_request(task: dict[str, Any], request: ExportRequest) -> bool:
    if str(task.get("type") or "") != EXPORT_TYPE or str(task.get("name") or "") != EXPORT_NAME:
        return False
    raw_where = task.get("param_where")
    try:
        where = json.loads(raw_where) if isinstance(raw_where, str) else raw_where
    except json.JSONDecodeError:
        return False
    warehouse = str(request.warehouse_id)
    values = list(_walk_values(where))
    for key, value in values:
        if key == "warehouse_id":
            candidates = value if isinstance(value, list) else [value]
            if warehouse in {str(item) for item in candidates}:
                return True
    # Captured responses encode filters as ["warehouse_id", "in", ["26"]].
    serialized = json.dumps(where, ensure_ascii=False, separators=(",", ":"))
    return "warehouse_id" in serialized and f'"{warehouse}"' in serialized


def validate_workbook(path: Path, *, warehouse_code: str) -> int:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - expose the real parser failure
        raise YacangError("校验 XLSX", f"{type(exc).__name__}: {exc}") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise YacangError("校验 XLSX", f"工作表数量应为 1，实际为 {len(workbook.sheetnames)}")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(str(value or "").strip() for value in next(rows, ()))
        if headers != EXPECTED_HEADERS:
            raise YacangError("校验 XLSX", f"表头不匹配: {safe_remote_detail(headers)}")
        row_count = 0
        wrong_warehouses: set[str] = set()
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            actual = str(row[2] or "").strip() if len(row) > 2 else ""
            if actual != warehouse_code and len(wrong_warehouses) < 5:
                wrong_warehouses.add(actual or "[empty]")
        if wrong_warehouses:
            raise YacangError(
                "校验 XLSX",
                f"期望仓库 {warehouse_code}，文件包含其他仓库: {sorted(wrong_warehouses)}",
            )
        return row_count
    finally:
        workbook.close()


class YacangExportQueue:
    def __init__(self, client: YacangClient, *, clock: Callable[[], float] = time.monotonic) -> None:
        self.client = client
        self.clock = clock

    def wait_for_file(
        self,
        request: ExportRequest,
        *,
        baseline_ids: set[str],
        timeout_seconds: float,
        poll_interval_seconds: float,
        sleep: Callable[[float], None] = time.sleep,
    ) -> str:
        deadline = self.clock() + timeout_seconds
        last_candidates: list[dict[str, Any]] = []
        while True:
            tasks = self.client.list_downloads()
            last_candidates = [
                task for task in tasks
                if str(task.get("id") or "") not in baseline_ids and task_matches_request(task, request)
            ]
            for task in last_candidates:
                path = str(task.get("path") or "").strip()
                if path:
                    return path
            if self.clock() >= deadline:
                detail = safe_remote_detail(last_candidates or {"new_matching_tasks": 0})
                raise YacangError("等待导出队列", f"{timeout_seconds:g} 秒超时，最后状态: {detail}")
            sleep(poll_interval_seconds)


class YacangRiskController:
    """Fail closed on duplicate/unknown submissions and space production writes."""

    def __init__(
        self,
        *,
        minimum_submission_interval_seconds: float = 1.0,
        clock: Callable[[], float] = time.monotonic,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self.minimum_submission_interval_seconds = minimum_submission_interval_seconds
        self.clock = clock
        self.sleep = sleep
        self._submitted: set[int] = set()
        self._last_submission_at: float | None = None

    def before_submission(self, request: ExportRequest) -> None:
        expected = dict(WAREHOUSES).get(request.warehouse_code)
        if expected != request.warehouse_id:
            raise YacangError("风控检查", f"未授权仓库映射: {request.warehouse_code}/{request.warehouse_id}")
        if request.warehouse_id in self._submitted:
            raise YacangError("风控检查", f"本次运行已提交仓库: {request.warehouse_code}")
        if self._last_submission_at is not None:
            remaining = self.minimum_submission_interval_seconds - (self.clock() - self._last_submission_at)
            if remaining > 0:
                self.sleep(remaining)
        self._submitted.add(request.warehouse_id)
        self._last_submission_at = self.clock()


def _cached_file(path: Path, *, warehouse_code: str, max_age_seconds: float) -> int | None:
    if max_age_seconds <= 0 or not path.is_file():
        return None
    age = time.time() - path.stat().st_mtime
    if age < 0 or age > max_age_seconds:
        return None
    try:
        return validate_workbook(path, warehouse_code=warehouse_code)
    except YacangError:
        return None


def export_inventory_sales(
    *,
    start_date: Any,
    end_date: Any,
    mobile: str | None = None,
    password: str | None = None,
    output_dir: str | Path | None = None,
    timeout_seconds: float = DEFAULT_TIMEOUT_SECONDS,
    poll_interval_seconds: float = DEFAULT_POLL_INTERVAL_SECONDS,
    cache_max_age_seconds: float = DEFAULT_CACHE_MAX_AGE_SECONDS,
    client: YacangClient | None = None,
    sleep: Callable[[float], None] = time.sleep,
) -> dict[str, Any]:
    start, end = validate_date_range(start_date, end_date)
    account = str(mobile if mobile is not None else os.environ.get("LXE_YACANG_MOBILE", "")).strip()
    secret = str(password if password is not None else os.environ.get("LXE_YACANG_PASSWORD", "")).strip()
    if not account or not secret:
        raise YacangError("读取配置", "桌面设置中缺少雅仓账号或密码")
    if timeout_seconds <= 0 or poll_interval_seconds <= 0:
        raise ValueError("timeout_seconds 和 poll_interval_seconds 必须大于 0")

    destination_dir = Path(output_dir) if output_dir is not None else dataset_dir("yacang_inventory_sales")
    destination_dir.mkdir(parents=True, exist_ok=True)
    requests = [ExportRequest(code, warehouse_id, start, end) for code, warehouse_id in WAREHOUSES]
    results: list[dict[str, Any]] = []
    pending: list[tuple[ExportRequest, Path]] = []
    for request in requests:
        path = destination_dir / request.filename
        cached_rows = _cached_file(path, warehouse_code=request.warehouse_code, max_age_seconds=cache_max_age_seconds)
        if cached_rows is None:
            pending.append((request, path))
        else:
            results.append({
                "warehouse": request.warehouse_code,
                "warehouse_id": request.warehouse_id,
                "xlsx_path": str(path),
                "row_count": cached_rows,
                "source": "cache",
            })

    if pending:
        api = client or YacangClient()
        api.login(account, secret)
        queue = YacangExportQueue(api)
        risk = YacangRiskController(sleep=sleep)
        for request, path in pending:
            baseline_ids = {str(task.get("id") or "") for task in api.list_downloads()}
            risk.before_submission(request)
            api.create_inventory_sales_export(
                warehouse_id=request.warehouse_id,
                start_date=start,
                end_date=end,
            )
            remote_path = queue.wait_for_file(
                request,
                baseline_ids=baseline_ids,
                timeout_seconds=float(timeout_seconds),
                poll_interval_seconds=float(poll_interval_seconds),
                sleep=sleep,
            )
            api.download_xlsx(remote_path, path)
            try:
                row_count = validate_workbook(path, warehouse_code=request.warehouse_code)
            except Exception:
                path.unlink(missing_ok=True)
                raise
            results.append({
                "warehouse": request.warehouse_code,
                "warehouse_id": request.warehouse_id,
                "xlsx_path": str(path),
                "row_count": row_count,
                "source": "yacang",
            })

    order = {code: index for index, (code, _) in enumerate(WAREHOUSES)}
    results.sort(key=lambda item: order[str(item["warehouse"])])
    return {
        "success": True,
        "start_date": start,
        "end_date": end,
        "sales_window_days": 15,
        "warehouse_count": len(results),
        "exports": results,
        "xlsx_paths": [str(item["xlsx_path"]) for item in results],
    }


__all__ = [
    "DEFAULT_CACHE_MAX_AGE_SECONDS",
    "EXPECTED_HEADERS",
    "ExportRequest",
    "WAREHOUSES",
    "YacangExportQueue",
    "YacangRiskController",
    "export_inventory_sales",
    "task_matches_request",
    "validate_date_range",
    "validate_workbook",
]
