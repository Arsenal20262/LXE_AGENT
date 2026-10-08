from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest

from services.vietnam_replenishment import settings, workflow
from services.vietnam_replenishment.asset_contract import load_sku_parameters
from services.vietnam_replenishment.workbook import RecommendationConfig


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path))
    from shared import workspace
    monkeypatch.setattr(workspace, "_workspace_root", tmp_path)


def sku_map(path: Path, cost=10):
    book = Workbook()
    book.active.append(["SKU", "成本", "跨境价", "折扣价", "热销标记"])
    book.active.append(["VN-A", cost, 30, 25, 1])
    book.save(path)
    book.close()
    return path


def test_defaults_persist_and_reload_in_another_process():
    state = settings.read_state()
    assert state["parameters"] == settings.config_json(RecommendationConfig())
    assert state["sku_map"] is None
    raw = {**state["parameters"], "exchange_rate": "4100", "sales_weight_7d": "1", "sales_weight_15d": "0", "sales_weight_30d": "0"}
    settings.save_parameters(raw)
    child = subprocess.run([sys.executable, "-I", "-B", "-m", "services.vietnam_replenishment.settings", "read"], capture_output=True, text=True)
    assert child.returncode == 0, child.stderr
    assert json.loads(child.stdout)["data"]["parameters"] == raw


@pytest.mark.parametrize("change", [
    {"sales_weight_7d": "0.5"}, {"sales_weight_7d": "-0.1"}, {"exchange_rate": "0"},
    {"day_adjustment_15d": "-1"}, {"day_adjustment_7d": "NaN"}, {"exchange_rate": "1e9999999999"},
    {"exchange_rate": "1e-9999"}, {"exchange_rate": "0.1234567890123456"}, {"sales_weight_7d": 0.6},
    {"extra": "1"},
])
def test_invalid_save_preserves_previous_file(change):
    old = settings.read_state()["parameters"]
    file = settings.data_directory() / "parameters.json"
    before = file.read_bytes()
    with pytest.raises((ValueError, ArithmeticError)):
        settings.save_parameters({**old, **change})
    assert file.read_bytes() == before


def test_corrupt_existing_configuration_is_reported_not_reset():
    settings.read_state()
    file = settings.data_directory() / "parameters.json"
    file.write_text("{broken")
    state = settings.read_state()
    assert "JSONDecodeError" in state["parameters_error"]
    assert state["parameters"] is None
    assert file.read_text() == "{broken"


def test_upload_is_atomic_and_failed_replacement_keeps_file(tmp_path):
    source = sku_map(tmp_path / "first.xlsx")
    state = settings.upload_map(source)
    target = Path(state["sku_map"]["path"])
    before = target.read_bytes()
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(RuntimeError, match="ZIP"):
        settings.upload_map(bad)
    assert target.read_bytes() == before
    state = settings.upload_map(sku_map(tmp_path / "second.xlsx", cost=12))
    assert load_sku_parameters(target)["VN-A"].cost == Decimal(12)
    assert state["sku_map"]["updated_at"]
    assert sorted(path.name for path in target.parent.glob("*.xlsx")) == ["sku-map.xlsx"]


@pytest.mark.parametrize(
    ("column", "label", "invalid"),
    [(column, label, invalid) for column, label in ((2, "成本"), (3, "跨境价"), (4, "折扣价"))
     for invalid in (None, 0)] + [(5, "热销标记", invalid) for invalid in (None, 0, 3)],
)
def test_incomplete_replacement_preserves_saved_data_and_reports_invalid_saved_map(
    tmp_path, column, label, invalid,
):
    state = settings.upload_map(sku_map(tmp_path / "valid.xlsx"))
    target = Path(state["sku_map"]["path"])
    parameters = settings.data_directory() / "parameters.json"
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (target, parameters)}
    replacement = sku_map(tmp_path / "不完整 映射.xlsx")
    book = load_workbook(replacement)
    book.active.cell(2, column).value = invalid
    book.save(replacement)
    book.close()
    location = "Sheet!E2" if column == 5 and invalid in (0, 3) else "VN-A"

    with pytest.raises(RuntimeError, match=f"{location}.*{label}"):
        settings.upload_map(replacement)
    assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in before} == before
    assert settings.read_state()["sku_map"]["updated_at"] == state["sku_map"]["updated_at"]

    # Files edited manually or saved by an earlier development build are checked,
    # never silently completed or given a default flag.
    target.write_bytes(replacement.read_bytes())
    invalid_bytes = target.read_bytes()
    state = settings.read_state()
    assert state["sku_map"] is None
    assert location in state["sku_map_error"] and label in state["sku_map_error"]
    export = tmp_path / "existing.xlsx"
    export.write_bytes(b"preserve destination")
    with pytest.raises(RuntimeError, match=f"{location}.*{label}"):
        settings.export_map("current", export)
    assert export.read_bytes() == b"preserve destination"
    assert target.read_bytes() == invalid_bytes
    assert (parameters.read_bytes(), parameters.stat().st_mtime_ns) == before[parameters]


def test_explicit_file_overrides_saved_without_changing_it_and_snapshot_is_stable(tmp_path):
    settings.upload_map(sku_map(tmp_path / "saved.xlsx", cost=10))
    explicit = sku_map(tmp_path / "explicit.xlsx", cost=20)
    with settings.run_inputs("explicit.xlsx") as (snapshot, config, _, source):
        assert source == str(explicit.resolve())
        assert load_sku_parameters(snapshot)["VN-A"].cost == Decimal(20)
        settings.save_parameters(settings.config_json(replace(config, exchange_rate=Decimal(4200))))
        settings.upload_map(sku_map(tmp_path / "replacement.xlsx", cost=30))
        explicit.write_bytes(b"changed after snapshot")
        assert load_sku_parameters(snapshot)["VN-A"].cost == Decimal(20)
        assert config.exchange_rate == Decimal(3900)
    assert not snapshot.exists()
    with settings.run_inputs() as (snapshot, config, _, source):
        assert load_sku_parameters(snapshot)["VN-A"].cost == Decimal(30)
        assert config.exchange_rate == Decimal(4200)
        assert source.endswith("sku-map.xlsx")


def test_bad_explicit_path_does_not_fall_back_or_start_erp(tmp_path, monkeypatch):
    settings.upload_map(sku_map(tmp_path / "valid.xlsx"))
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: pytest.fail("Source loading must not run"))
    with pytest.raises(workflow.VietnamWorkflowError, match="missing.xlsx"):
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx", sku_map=str(tmp_path / "missing.xlsx"))


def test_invalid_configuration_stops_before_erp(tmp_path, monkeypatch):
    settings.upload_map(sku_map(tmp_path / "valid.xlsx"))
    path = settings.data_directory() / "parameters.json"
    path.write_text('{"exchange_rate":"0"}')
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: pytest.fail("Source loading must not run"))
    with pytest.raises(ValueError):
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")



@pytest.mark.parametrize("old_parameters", ["{broken", json.dumps({"vietnam_recommendation": {"exchange_rate": "4200"}})])
def test_initialization_uses_only_skill_data_directory(tmp_path, old_parameters):
    app = tmp_path / "app"
    (app / "config").mkdir(parents=True)
    (app / "config" / "settings.json").write_text(old_parameters)
    legacy_maps = app / "inputs" / "vietnam" / "sku_parameter_map"
    legacy_maps.mkdir(parents=True)
    (legacy_maps / "manifest.json").write_text("{broken")
    sku_map(legacy_maps / "old.xlsx")
    state = settings.read_state()
    assert state["parameters"] == settings.config_json(RecommendationConfig())
    assert state["parameters_error"] is None
    assert state["sku_map"] is None
    assert state["sku_map_error"] is None
    with pytest.raises(RuntimeError, match="请在越南备货设置上传"):
        with settings.run_inputs():
            pytest.fail("Old map must not be imported")
    assert (app / "config" / "settings.json").read_text() == old_parameters


def test_saved_map_corruption_is_reported_and_upload_can_repair(tmp_path):
    state = settings.upload_map(sku_map(tmp_path / "initial.xlsx"))
    target = Path(state["sku_map"]["path"])
    target.write_bytes(b"broken")
    state = settings.read_state()
    assert "ZIP" in state["sku_map_error"]
    assert state["sku_map"] is None
    assert target.read_bytes() == b"broken"
    state = settings.upload_map(sku_map(tmp_path / "replacement.xlsx"))
    assert state["sku_map_error"] is None
    assert load_sku_parameters(target)["VN-A"].cost == Decimal(10)


def test_desktop_upload_and_saved_map_read_return_all_errors_without_replacing_data(tmp_path):
    state = settings.upload_map(sku_map(tmp_path / "good.xlsx"))
    target = Path(state["sku_map"]["path"])
    parameters = settings.data_directory() / "parameters.json"
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (target, parameters)}
    invalid = sku_map(tmp_path / "中文 不完整.xlsx")
    book = load_workbook(invalid)
    book.active["B2"] = 0
    book.active["E2"] = None
    book.active.append(["VN-B", 12, "bad", 0, 1])
    book.save(invalid)
    book.close()
    child = subprocess.run(
        [sys.executable, "-I", "-B", "-m", "services.vietnam_replenishment.settings", "upload", str(invalid)],
        capture_output=True, text=True, encoding="utf-8",
    )
    assert child.returncode == 1, child.stderr
    result = json.loads(child.stdout)
    assert result["success"] is False
    assert "4 个错误，涉及 2 行" in result["error"]
    for cell in ("B2", "E2", "C3", "D3"):
        assert f"Sheet!{cell}" in result["error"]
    assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in before} == before
    assert not list(target.parent.glob(".sku-map-*.xlsx"))

    # A map edited outside the app uses exactly the same validation on read.
    target.write_bytes(invalid.read_bytes())
    state = settings.read_state()
    assert state["sku_map"] is None
    assert state["sku_map_error"] == result["error"]
    assert target.read_bytes() == invalid.read_bytes()
    assert (parameters.read_bytes(), parameters.stat().st_mtime_ns) == before[parameters]
