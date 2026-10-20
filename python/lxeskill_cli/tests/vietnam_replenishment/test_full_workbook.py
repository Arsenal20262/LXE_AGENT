"""The exported nine-sheet workbook retains the user's four business tabs."""

from __future__ import annotations

from hashlib import sha256
from importlib.resources import files
import os
from pathlib import Path
from xml.dom import minidom
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font
import pytest

from services.vietnam_replenishment.asset_contract import AssetContractError, load_template_sku_parameters
from services.vietnam_replenishment.full_workbook import (
    MAIN, CURRENT, REQUIRED_ORDER, UNCHANGED, FullWorkbookError, _sheet_parts,
    combine_full_workbook, validate_full_template,
)


def _inputs(directory: Path) -> tuple[Path, Path]:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as handle:
        original = load_workbook(handle)
    for index, name in enumerate(UNCHANGED, 1):
        sheet = original.create_sheet(name, index)
        sheet["A1"] = f"original {name}"
        sheet["A1"].font = Font(color="FF336699")
        sheet["B2"] = "=越南备货清单!H2"
        sheet["C3"] = "SKU-PREVIOUS"
    template = directory / "template.xlsx"
    original.save(template)
    original.close()

    with resource.open("rb") as handle:
        fresh = load_workbook(handle)
    fresh["越南备货清单"]["E2"] = "SKU-CURRENT"
    fresh["越南备货清单"]["E2"].font = Font(bold=True, color="FF112233")
    fresh["雅仓库存"]["B2"] = "SKU-CURRENT"
    current = directory / "current.xlsx"
    fresh.save(current)
    fresh.close()
    return template, current


def test_four_unchanged_sheets_keep_exact_xml_and_current_five_are_inserted(tmp_path: Path) -> None:
    template, current = _inputs(tmp_path)
    before = sha256(template.read_bytes()).digest()
    output = tmp_path / "final.xlsx"

    combine_full_workbook(template, current, output)

    assert sha256(template.read_bytes()).digest() == before
    with ZipFile(template) as original, ZipFile(output) as result:
        names = _sheet_parts(original)
        assert result.testzip() is None
        for name in UNCHANGED:
            assert result.read(names[name]) == original.read(names[name])

    book = load_workbook(output, read_only=False, data_only=False)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        assert book["越南备货清单"]["E2"].value == "SKU-CURRENT"
        assert book["越南备货清单"]["E2"].font.bold
        assert book["雅仓库存"]["B2"].value == "SKU-CURRENT"
        for name in UNCHANGED:
            assert book[name]["A1"].value == f"original {name}"
            assert book[name]["A1"].font.color.rgb == "FF336699"
            assert book[name]["B2"].value == "=越南备货清单!H2"
            assert book[name]["C3"].value == "SKU-PREVIOUS"
    finally:
        book.close()


def test_bad_sheet_order_is_rejected_without_partial_output(tmp_path: Path) -> None:
    template, current = _inputs(tmp_path)
    book = load_workbook(template)
    book.move_sheet("改库存", offset=1)
    book.save(template)
    book.close()
    output = tmp_path / "final.xlsx"

    with pytest.raises(FullWorkbookError, match="指定顺序的九表"):
        combine_full_workbook(template, current, output)
    assert not output.exists()


def _builtin_template(path: Path) -> Path:
    """A valid nine-sheet workbook with built-in formats and no optional numFmts."""
    book = Workbook()
    book.remove(book.active)
    for name in REQUIRED_ORDER:
        book.create_sheet(name)
    main = book[REQUIRED_ORDER[0]]
    for cell, value in {"B1": "判断热销", "E1": "SKU", "G1": "成本", "AE1": "跨境价", "AJ1": "折扣价",
                        "B2": 2, "E2": "VN-A", "G2": 12, "AE2": 22, "AJ2": 16}.items():
        main[cell] = value
    for name in UNCHANGED:
        book[name]["A1"] = f"original {name}"
        book[name]["A2"] = 1234.5
        book[name]["A2"].number_format = "0.00"
        book[name]["B2"] = "=越南备货清单!H2"
    book.save(path)
    book.close()
    staged = path.with_suffix(".tmp")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                styles = ET.fromstring(data)
                formats = styles.find(f"{{{MAIN}}}numFmts")
                assert formats is not None and len(formats) == 0
                styles.remove(formats)
                data = ET.tostring(styles)
            target.writestr(item, data)
    staged.replace(path)
    return path


def _style_namespace(path: Path, prefix: str | None, section_default: str | None = None) -> None:
    """Use an explicit SpreadsheetML binding without changing style semantics."""
    staged = path.with_suffix(".tmp")
    with ZipFile(path) as source, ZipFile(staged, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                document = minidom.parseString(data)
                try:
                    for node in list(document.getElementsByTagNameNS(MAIN, "*")):
                        document.renameNode(node, MAIN, f"{prefix}:{node.localName}" if prefix else node.localName)
                    root = document.documentElement
                    for attribute in list(root.attributes.values()):
                        if attribute.value == MAIN and attribute.name.startswith("xmlns"):
                            root.removeAttribute(attribute.name)
                    root.setAttribute(f"xmlns:{prefix}" if prefix else "xmlns", MAIN)
                    if section_default is not None:
                        root.setAttribute("xmlns", MAIN)
                        for name in ("fonts", "cellXfs"):
                            document.getElementsByTagNameNS(MAIN, name)[0].setAttribute("xmlns", section_default)
                    data = document.toxml(encoding="utf-8")
                finally:
                    document.unlink()
                assert all(node.tag.startswith(f"{{{MAIN}}}") for node in ET.fromstring(data).iter())
            target.writestr(item, data)
    staged.replace(path)


@pytest.mark.parametrize("template_prefix,section_default", [
    (None, None), ("ns0", None), ("sheet", "urn:unused-style-default"),
])
@pytest.mark.parametrize("calculated_prefix", [None, "calc"])
@pytest.mark.parametrize("missing_formats", [False, True])
def test_style_merge_accepts_optional_numfmts_and_keeps_all_style_references_valid(
    tmp_path, missing_formats, template_prefix, section_default, calculated_prefix,
):
    template, current = _inputs(tmp_path)
    if missing_formats:
        _builtin_template(template)
    book = load_workbook(current)
    book[CURRENT[0]].conditional_formatting.add(
        "E2", CellIsRule(operator="equal", formula=['"SKU-CURRENT"'], font=Font(bold=True)),
    )
    book.save(current)
    book.close()
    _style_namespace(template, template_prefix, section_default)
    _style_namespace(current, calculated_prefix)
    validate_full_template(template)
    output = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, output)

    with ZipFile(template) as original, ZipFile(output) as merged:
        before = ET.fromstring(original.read("xl/styles.xml"))
        after = ET.fromstring(merged.read("xl/styles.xml"))
        assert all(node.tag.startswith(f"{{{MAIN}}}") for node in after.iter())
        assert [node.tag for node in after][:2] == [f"{{{MAIN}}}numFmts", f"{{{MAIN}}}fonts"]
        assert [node.tag for node in after].index(f"{{{MAIN}}}dxfs") < [node.tag for node in after].index(f"{{{MAIN}}}tableStyles")
        formats = after.find(f"{{{MAIN}}}numFmts")
        assert formats is not None and all(node.tag == f"{{{MAIN}}}numFmt" for node in formats)
        identifiers = [int(node.attrib["numFmtId"]) for node in formats]
        assert len(identifiers) == len(set(identifiers))
        assert int(formats.attrib["count"]) == len(identifiers)
        for section in ("cellStyleXfs", "cellXfs"):
            old = before.find(f"{{{MAIN}}}{section}")
            new = after.find(f"{{{MAIN}}}{section}")
            assert [ET.tostring(node) for node in new[:len(old)]] == [ET.tostring(node) for node in old]
            for node in new:
                assert node.tag == f"{{{MAIN}}}xf"
                number = int(node.attrib.get("numFmtId", 0))
                assert number < 164 or number in identifiers
                for attribute, name in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders"), ("xfId", "cellStyleXfs")):
                    if attribute in node.attrib:
                        assert int(node.attrib[attribute]) < len(after.find(f"{{{MAIN}}}{name}"))
        parts = _sheet_parts(original)
        assert all(original.read(parts[name]) == merged.read(parts[name]) for name in UNCHANGED)

    result, reference = load_workbook(output), load_workbook(current)
    try:
        assert tuple(result.sheetnames) == REQUIRED_ORDER
        for name in CURRENT:
            assert result[name].max_row == reference[name].max_row
            assert result[name].max_column == reference[name].max_column
            for row in reference[name]:
                for cell in row:
                    assert result[name][cell.coordinate].number_format == cell.number_format
        if missing_formats:
            assert result[UNCHANGED[0]]["A2"].number_format == "0.00"
    finally:
        result.close()
        reference.close()


@pytest.mark.parametrize("template_prefix,calculated_prefix,missing_formats", [
    (None, None, False), ("ns0", None, True),
    ("sheet", "calc", False), ("sheet", "calc", True),
])
@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_template_styles_can_be_reopened_by_office(tmp_path, template_prefix, calculated_prefix, missing_formats):
    from services.vietnam_replenishment.recalculation import recalculate_with_office

    template, current = _inputs(tmp_path)
    if missing_formats:
        _builtin_template(template)
    _style_namespace(template, template_prefix)
    _style_namespace(current, calculated_prefix)
    validate_full_template(template)
    merged = tmp_path / "merged.xlsx"
    combine_full_workbook(template, current, merged)
    reopened = tmp_path / "office.xlsx"
    recalculate_with_office(merged, reopened)
    book = load_workbook(reopened, data_only=True)
    try:
        assert tuple(book.sheetnames) == REQUIRED_ORDER
        assert book[REQUIRED_ORDER[0]]["E2"].value == "SKU-CURRENT"
        if missing_formats:
            for name in UNCHANGED:
                assert book[name]["A2"].value == 1234.5
                assert book[name]["A2"].number_format == "0.00"
    finally:
        book.close()


def test_template_parameters_use_saved_formula_result_without_a_sku_map(tmp_path: Path) -> None:
    template, _ = _inputs(tmp_path)
    book = load_workbook(template)
    main = book["越南备货清单"]
    for cell, value in {"B2": 2, "E2": "SKU-CURRENT", "G2": 10,
                        "AE2": "=G2*20", "AJ2": None}.items():
        main[cell] = value
    book.save(template)
    book.close()
    cached = tmp_path / "cached.xlsx"
    with ZipFile(template) as source, ZipFile(cached, "w") as target:
        part = _sheet_parts(source)["越南备货清单"]
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == part:
                xml = ET.fromstring(payload)
                cell = xml.find(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c[@r='AE2']")
                cell.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text = "200"
                payload = ET.tostring(xml)
            target.writestr(info, payload)
    actual = load_template_sku_parameters(cached, {"SKU-CURRENT", "SKU-NEW"})
    assert actual["SKU-CURRENT"].cost == 10
    assert actual["SKU-CURRENT"].cross_border_price == 200
    assert actual["SKU-CURRENT"].discount_price is None
    assert "SKU-NEW" not in actual


def test_template_conflicting_duplicate_current_sku_is_rejected(tmp_path: Path) -> None:
    template, _ = _inputs(tmp_path)
    book = load_workbook(template)
    main = book["越南备货清单"]
    for row, cost in ((2, 10), (3, 11)):
        for column, value in (("B", 2), ("E", "SKU-CURRENT"), ("G", cost),
                              ("AE", 20), ("AJ", 15)):
            main[f"{column}{row}"] = value
    book.save(template)
    book.close()
    with pytest.raises(AssetContractError, match="参数冲突"):
        load_template_sku_parameters(template, {"SKU-CURRENT"})
