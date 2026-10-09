from pathlib import Path

from openpyxl import Workbook
import pytest

from services.yacang.validation import (
    INVENTORY_LIST_HEADERS,
    INVENTORY_SALES_DELIVERY_HEADERS,
    INVENTORY_SALES_HEADERS,
    WAREHOUSE_PRODUCTS_HEADERS,
)
from services.vietnam_replenishment import yacang_sources


def _workbook(path: Path, headers: tuple[str, ...], rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for values in rows:
        sheet.append([values.get(header) for header in headers])
    workbook.save(path)
    workbook.close()


def _artifacts(tmp_path: Path, *, sales=None, inventory=None, products=None):
    reports = (
        ("inventory-sales", INVENTORY_SALES_HEADERS, sales or [], "VN8806"),
        ("inventory-current-snapshot", INVENTORY_LIST_HEADERS, inventory or [], "VN8806"),
        ("warehouse-products", WAREHOUSE_PRODUCTS_HEADERS, products or [], None),
    )
    artifacts = []
    for report, headers, rows, warehouse in reports:
        path = tmp_path / f"{report}.xlsx"
        _workbook(path, headers, rows)
        artifacts.append({"report": report, "warehouse": warehouse, "created_date": None, "path": str(path)})
    return artifacts


def test_union_and_global_product_filter_keep_missing_source_diagnostics(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[
            {"SKU": " VN-A ", "仓库": "VN8806", "7天销量": 7, "在途": 3},
            {"SKU": "VN-C", "仓库": "VN8806", "7天销量": 4},
        ],
        inventory=[
            {"SKU": "VN-A", "仓库": "VN8806", "库存数量": 10, "在途数量": 2},
            {"SKU": "VN-B", "仓库": "VN8806", "库存数量": 5},
        ],
        products=[
            {"SKU": "VN-A", "创建时间": "2026-09-23 10:00"},
            {"SKU": "VN-B", "创建时间": "2026-09-23 11:00"},
            {"SKU": "OTHER-COUNTRY", "创建时间": "2026-09-23 12:00"},
        ],
    )

    sources = yacang_sources.load_vietnam_sources(artifacts)

    assert sources.skus == ("VN-A", "VN-B", "VN-C")
    assert sources.missing_sales == ("VN-B",)
    assert sources.missing_inventory == ("VN-C",)
    assert sources.missing_products == ("VN-C",)
    assert sources.sales["VN-A"]["7天销量"] == 7
    assert sources.inventory["VN-A"]["在途数量"] == 2
    assert sources.in_transit == {"VN-A": 2, "VN-B": None, "VN-C": None}
    assert sources.missing_in_transit == ("VN-B", "VN-C")
    assert sources.in_transit_mismatch == ("VN-A",)
    assert set(sources.products) == {"VN-A", "VN-B"}
    assert sources.products["VN-A"]["创建时间"] == "2026-09-23 10:00"
    assert "上架时间" not in sources.products["VN-A"]
    assert sources.artifacts["warehouse-products"] == Path(artifacts[2]["path"])


def test_inventory_row_with_blank_transit_is_reported_missing(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806", "在途": 8}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806", "在途数量": None}],
    )
    sources = yacang_sources.load_vietnam_sources(artifacts)
    assert sources.missing_inventory == ()
    assert sources.in_transit == {"VN-A": None}
    assert sources.missing_in_transit == ("VN-A",)
    assert sources.in_transit_mismatch == ()


def test_invalid_authoritative_in_transit_is_rejected(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806", "在途": 4}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806", "在途数量": "invalid"}],
    )
    with pytest.raises(yacang_sources.VietnamSourceError, match="在途数量"):
        yacang_sources.load_vietnam_sources(artifacts)


def test_invalid_sales_transit_is_not_silently_ignored(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806", "在途": "invalid"}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806", "在途数量": 3}],
    )
    with pytest.raises(yacang_sources.VietnamSourceError, match="inventory-sales SKU VN-A 的 在途"):
        yacang_sources.load_vietnam_sources(artifacts)


def test_formula_in_retained_source_row_is_rejected(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806", "7天销量": "=2+2"}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806"}],
    )
    with pytest.raises(yacang_sources.VietnamSourceError, match="公式"):
        yacang_sources.load_vietnam_sources(artifacts)


@pytest.mark.parametrize("report,rows", [
    ("inventory-sales", [{"SKU": "", "仓库": "VN8806"}]),
    ("inventory-current-snapshot", [{"SKU": 123, "仓库": "VN8806"}]),
    ("warehouse-products", [{"SKU": None, "创建时间": "2026-09-23 10:00"}]),
])
def test_reject_blank_or_nontext_sku(tmp_path, report, rows):
    data = {
        "sales": [{"SKU": "VN-A", "仓库": "VN8806"}],
        "inventory": [{"SKU": "VN-A", "仓库": "VN8806"}],
        "products": [{"SKU": "VN-A", "创建时间": "2026-09-23 10:00"}],
    }
    data[{"inventory-sales": "sales", "inventory-current-snapshot": "inventory", "warehouse-products": "products"}[report]] = rows
    with pytest.raises(yacang_sources.VietnamSourceError, match="SKU"):
        yacang_sources.load_vietnam_sources(_artifacts(tmp_path, **data))


def test_reject_duplicate_sku_and_wrong_warehouse(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806"}, {"SKU": " VN-A ", "仓库": "VN8806"}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806"}],
    )
    with pytest.raises(yacang_sources.VietnamSourceError, match="重复 SKU"):
        yacang_sources.load_vietnam_sources(artifacts)

    artifacts[0]["warehouse"] = "MY8801"
    with pytest.raises(yacang_sources.VietnamSourceError, match="VN8806"):
        yacang_sources.load_vietnam_sources(artifacts)


def test_reject_incomplete_or_filtered_artifact_set(tmp_path):
    artifacts = _artifacts(tmp_path, sales=[{"SKU": "VN-A", "仓库": "VN8806"}])
    with pytest.raises(yacang_sources.VietnamSourceError, match="三份"):
        yacang_sources.load_vietnam_sources(artifacts[:2])
    artifacts[0]["created_date"] = {"start_date": "2026-09-01", "end_date": "2026-09-30"}
    with pytest.raises(yacang_sources.VietnamSourceError, match="创建日期"):
        yacang_sources.load_vietnam_sources(artifacts)


def test_empty_current_vietnam_reports_do_not_produce_an_operator_sku_set(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        products=[{"SKU": "GLOBAL-ONLY", "创建时间": "2026-09-23 10:00"}],
    )

    with pytest.raises(yacang_sources.VietnamSourceError, match="均无 SKU"):
        yacang_sources.load_vietnam_sources(artifacts)



def _valid_file_arguments(tmp_path):
    artifacts = _artifacts(
        tmp_path,
        sales=[{"SKU": "VN-A", "仓库": "VN8806", "7天销量": 7, "15天销量": 15, "30天销量": 30, "在途": 5}],
        inventory=[{"SKU": "VN-A", "仓库": "VN8806", "可用库存": 10, "在途数量": 3}],
        products=[{"SKU": "VN-A", "中文标题": "测试产品", "创建时间": "2026-09-23 10:00"}],
    )
    return dict(zip(("sales", "inventory", "products"), (entry["path"] for entry in artifacts)))


def test_local_files_resolve_workspace_paths_and_record_snapshot_hashes(tmp_path, monkeypatch):
    import hashlib
    from shared import workspace

    paths = _valid_file_arguments(tmp_path)
    original = {name: Path(path).read_bytes() for name, path in paths.items()}
    monkeypatch.setattr(workspace, "_workspace_root", tmp_path)
    sources = yacang_sources.load_vietnam_files(**{name: Path(path).name for name, path in paths.items()})
    assert sources.skus == ("VN-A",)
    assert sources.in_transit["VN-A"] == 3
    assert sources.in_transit_mismatch == ("VN-A",)
    for report, name in zip(("inventory-sales", "inventory-current-snapshot", "warehouse-products"), paths):
        assert sources.artifacts[report] == Path(paths[name])
        assert sources.file_hashes[report] == hashlib.sha256(original[name]).hexdigest()
        assert Path(paths[name]).read_bytes() == original[name]


def test_delivered_sales_without_date_preserve_calculation_inputs(tmp_path):
    from services.yacang.delivery import remove_inventory_sales_creation_date

    paths = _valid_file_arguments(tmp_path)
    original = yacang_sources.load_vietnam_files(**paths)
    sales_path = Path(paths["sales"])
    original_bytes = sales_path.read_bytes()
    delivered = tmp_path / "库存动销 交付.xlsx"
    delivered.write_bytes(original_bytes)
    remove_inventory_sales_creation_date(delivered)
    paths["sales"] = str(delivered)

    actual = yacang_sources.load_vietnam_files(**paths)
    for field in (
        "skus", "sales", "inventory", "products", "in_transit", "missing_sales",
        "missing_inventory", "missing_products", "missing_in_transit", "in_transit_mismatch",
    ):
        assert getattr(actual, field) == getattr(original, field)
    assert actual.artifacts["inventory-sales"] == delivered
    assert actual.file_hashes["inventory-sales"] != original.file_hashes["inventory-sales"]
    assert sales_path.read_bytes() == original_bytes


@pytest.mark.parametrize("problem", ["missing_column", "extra_column", "wrong_warehouse"])
def test_delivered_sales_still_reject_invalid_schema_and_warehouse(tmp_path, problem):
    paths = _valid_file_arguments(tmp_path)
    headers = INVENTORY_SALES_DELIVERY_HEADERS
    warehouse = "VN8806"
    if problem == "missing_column":
        headers = tuple(header for header in headers if header != "7天销量")
    elif problem == "extra_column":
        headers = (*headers, "额外字段")
    else:
        warehouse = "MY8801"
    _workbook(Path(paths["sales"]), headers, [{"SKU": "VN-A", "仓库": warehouse}])
    expected = "MY8801" if problem == "wrong_warehouse" else "表头不匹配"
    with pytest.raises(yacang_sources.VietnamSourceError, match=expected):
        yacang_sources.load_vietnam_files(**paths)


def test_local_reports_are_frozen_before_validation_and_temp_files_removed(tmp_path, monkeypatch):
    paths = _valid_file_arguments(tmp_path)
    original_reader = yacang_sources._read_rows
    seen = []

    def read_snapshot(path, report, **kwargs):
        seen.append(path)
        if len(seen) == 1:
            for original in paths.values():
                Path(original).write_bytes(b"changed during calculation")
        return original_reader(path, report, **kwargs)

    monkeypatch.setattr(yacang_sources, "_read_rows", read_snapshot)
    sources = yacang_sources.load_vietnam_files(**paths)
    assert sources.sales["VN-A"]["7天销量"] == 7
    assert sources.inventory["VN-A"]["可用库存"] == 10
    assert len(seen) == 3
    assert all(not path.exists() for path in seen)


@pytest.mark.parametrize("problem", ["missing", "directory", "corrupt", "wrong_report", "wrong_warehouse"])
def test_invalid_local_report_stops_with_observed_diagnostic(tmp_path, problem):
    from openpyxl import load_workbook

    paths = _valid_file_arguments(tmp_path)
    sales_path = Path(paths["sales"])
    if problem == "missing":
        sales_path.unlink()
    elif problem == "directory":
        sales_path.unlink()
        sales_path.mkdir()
    elif problem == "corrupt":
        sales_path.write_bytes(b"not a zip")
    elif problem == "wrong_report":
        sales_path.write_bytes(Path(paths["inventory"]).read_bytes())
    else:
        book = load_workbook(sales_path)
        book.active["C2"] = "MY8801"
        book.save(sales_path)
        book.close()
    expected = {
        "missing": "必须是存在的 XLSX", "directory": "必须是存在的 XLSX",
        "corrupt": "BadZipFile", "wrong_report": "表头不匹配", "wrong_warehouse": "MY8801",
    }[problem]
    with pytest.raises(yacang_sources.VietnamSourceError, match=expected):
        yacang_sources.load_vietnam_files(**paths)
