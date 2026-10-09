"""The Vietnam workflow uses explicit local reports and one settings snapshot."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from services.vietnam_replenishment import workflow
from services.vietnam_replenishment.asset_contract import load_sku_parameters
from services.vietnam_replenishment.workbook import RecommendationConfig
from services.vietnam_replenishment.yacang_sources import VietnamSourceError, VietnamSources
from services.vietnam_replenishment import settings as store


def _sku_map(
    path: Path, *, sku: str | None = "VN-A", cost: int | None = 10,
    price: int | None = 20, discount: int | None = 15, hot_flag: int | None = 2,
) -> Path:
    book = Workbook()
    try:
        book.active.append(("SKU", "成本", "跨境价", "折扣价", "热销标记"))
        if sku is not None:
            book.active.append((sku, cost, price, discount, hot_flag))
        book.save(path)
    finally:
        book.close()
    return path


@pytest.fixture(autouse=True)
def isolated_maps(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path))
    for name in (
        "LXE_VIETNAM_WEIGHT_30D", "LXE_VIETNAM_WEIGHT_15D",
        "LXE_VIETNAM_WEIGHT_7D", "LXE_VIETNAM_EXCHANGE_RATE",
    ):
        monkeypatch.delenv(name, raising=False)


def _current(monkeypatch: pytest.MonkeyPatch, path: Path) -> Path:
    store.upload_map(path)
    return store.data_directory() / "sku-map.xlsx"


def _no_sources(monkeypatch: pytest.MonkeyPatch) -> None:
    def unexpected_load(**_paths) -> None:
        raise AssertionError("Report loading must not start before map preflight")

    monkeypatch.setattr(workflow, "load_vietnam_files", unexpected_load)


def _sources() -> VietnamSources:
    sku = "VN-A"
    return VietnamSources(
        skus=(sku,),
        sales={sku: {
            "SKU": sku, "仓库": "VN8806", "7天销量": 0,
            "15天销量": 0, "30天销量": 0, "在途": 99,
        }},
        inventory={sku: {
            "SKU": sku, "仓库": "VN8806", "库存数量": 7,
            "占用数量": 2, "在途数量": 3, "可用库存": 5,
        }},
        products={sku: {
            "SKU": sku, "中文标题": "合成产品 A", "创建时间": "2026-09-23 10:00",
        }},
        missing_sales=(), missing_inventory=(), missing_products=(),
        in_transit={sku: Decimal("3")}, missing_in_transit=(),
        in_transit_mismatch=(sku,), artifacts={},
    )


def test_missing_current_prompts_upload_before_loading_reports(monkeypatch: pytest.MonkeyPatch) -> None:
    _no_sources(monkeypatch)

    with pytest.raises(workflow.VietnamWorkflowError, match="上传") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "sku_parameter_map_invalid"


def test_legacy_current_without_manifest_stops_before_loading_reports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    legacy = tmp_path / "inputs" / "vietnam" / "sku_parameter_map" / "current"
    legacy.mkdir(parents=True)
    _sku_map(legacy / "untrusted.xlsx")
    _no_sources(monkeypatch)
    with pytest.raises(workflow.VietnamWorkflowError, match="上传") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "sku_parameter_map_invalid"


def test_broken_current_stops_before_loading_reports(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not an xlsx")
    version = _current(monkeypatch, _sku_map(tmp_path / "valid.xlsx"))
    version.write_bytes(path.read_bytes())
    _no_sources(monkeypatch)

    with pytest.raises(workflow.VietnamWorkflowError, match="ZIP") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "sku_parameter_map_invalid"


def test_empty_current_stops_before_loading_reports(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    invalid = _sku_map(tmp_path / "empty.xlsx", sku=None)
    version = _current(monkeypatch, _sku_map(tmp_path / "valid.xlsx"))
    version.write_bytes(invalid.read_bytes())
    _no_sources(monkeypatch)

    with pytest.raises(workflow.VietnamWorkflowError, match="没有 SKU") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "sku_parameter_map_invalid"


@pytest.mark.parametrize(
    ("label", "values"),
    [(label, {field: invalid})
     for field, label in (("cost", "成本"), ("price", "跨境价"), ("discount", "折扣价"))
     for invalid in (0, -1)] + [("热销标记", {"hot_flag": None})],
)
@pytest.mark.parametrize("explicit", (False, True))
def test_invalid_map_stops_before_loading_reports_without_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, label: str, values: dict, explicit: bool,
) -> None:
    saved = _current(monkeypatch, _sku_map(tmp_path / "valid.xlsx"))
    invalid = _sku_map(tmp_path / "不完整 映射.xlsx", **values)
    if not explicit:
        saved.write_bytes(invalid.read_bytes())
    before = saved.read_bytes()
    _no_sources(monkeypatch)
    with pytest.raises(workflow.VietnamWorkflowError, match=f"VN-A.*{label}") as error:
        workflow.generate_current_vietnam_recommendation(
            sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx",
            sku_map=str(invalid) if explicit else None,
        )
    assert error.value.code == "sku_parameter_map_invalid"
    assert saved.read_bytes() == before


def test_snapshot_is_private_stable_and_removed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    current = _sku_map(tmp_path / "current.xlsx", price=20)
    _current(monkeypatch, current)

    with store.run_inputs() as (snapshot, _config, _config_path, _map_path):
        assert snapshot != current
        assert snapshot.is_file()
        store.upload_map(_sku_map(tmp_path / "replacement.xlsx", price=40))
        assert load_sku_parameters(snapshot)["VN-A"].cross_border_price == Decimal("20")
    assert not snapshot.exists()


def test_local_reports_uses_snapshot_and_default_config(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    current = _sku_map(tmp_path / "current.xlsx", price=20)
    _current(monkeypatch, current)
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    sources = _sources()
    calls: list[str] = []

    def fake_load(**_paths) -> VietnamSources:
        calls.append("load")
        store.upload_map(_sku_map(tmp_path / "replacement.xlsx", price=40))
        return sources

    def fake_generate(map_path: Path, output_path: Path, *, sources: VietnamSources, config: RecommendationConfig) -> Path:
        calls.append("generate")
        assert map_path != current
        assert load_sku_parameters(map_path)["VN-A"].cross_border_price == Decimal("20")
        assert sources is expected_sources
        assert config == RecommendationConfig()
        output_path.parent.mkdir(parents=True)
        output_path.write_bytes(b"synthetic final workbook")
        return output_path

    expected_sources = sources
    monkeypatch.setattr(workflow, "load_vietnam_files", fake_load)
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fake_generate)

    result = workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert calls == ["load", "generate"]
    assert result.sku_count == 1
    assert result.config == RecommendationConfig()
    assert result.config_source == str(store.data_directory() / "parameters.json")
    assert result.output_xlsx.name == "越南备货清单.xlsx"
    assert result.output_xlsx.read_bytes() == b"synthetic final workbook"




def test_nondefault_config_reaches_generator_and_result(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    store.save_parameters(store.config_json(RecommendationConfig(Decimal("0.7"), Decimal("0.6"), Decimal("0.1"), Decimal("4000"))))
    calls: list[str] = []

    def fake_load(**_paths) -> VietnamSources:
        calls.append("load")
        return _sources()

    def fake_generate(
        _map_path: Path, output_path: Path, *, sources: VietnamSources, config: RecommendationConfig,
    ) -> Path:
        calls.append("generate")
        assert sources.skus == ("VN-A",)
        assert config == RecommendationConfig(
            day_adjustment_30d=Decimal("0.7"), day_adjustment_15d=Decimal("0.6"),
            day_adjustment_7d=Decimal("0.1"), exchange_rate=Decimal("4000"),
        )
        output_path.parent.mkdir(parents=True)
        output_path.write_bytes(b"synthetic final workbook")
        return output_path

    monkeypatch.setattr(workflow, "load_vietnam_files", fake_load)
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fake_generate)
    result = workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert calls == ["load", "generate"]
    assert result.config.exchange_rate == Decimal("4000")
    assert result.config_source == str(store.data_directory() / "parameters.json")




def test_source_error_does_not_start_generation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    output_root = tmp_path / "artifacts"
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: output_root / parts[-1])

    def fail_load(**_paths) -> VietnamSources:
        raise VietnamSourceError("三份来源中库存列表失败")

    monkeypatch.setattr(workflow, "load_vietnam_files", fail_load)
    monkeypatch.setattr(
        workflow, "generate_vietnam_workbook",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("generator called")),
    )

    with pytest.raises(VietnamSourceError, match="库存列表失败"):
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert not output_root.exists()


def test_empty_source_sku_set_does_not_start_generation(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    output_root = tmp_path / "artifacts"
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: output_root / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: replace(_sources(), skus=()))
    monkeypatch.setattr(
        workflow, "generate_vietnam_workbook",
        lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("generator called")),
    )

    with pytest.raises(workflow.VietnamWorkflowError, match="没有 SKU") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "current_skus_empty"
    assert not output_root.exists()


def test_failed_generation_removes_partial_final_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    output_root = tmp_path / "artifacts"
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: output_root / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())

    def fail_generate(_map_path: Path, output_path: Path, *, sources: VietnamSources, config: RecommendationConfig) -> Path:
        output_path.parent.mkdir(parents=True)
        output_path.write_bytes(b"incomplete")
        raise RuntimeError("Office actual failure")

    monkeypatch.setattr(workflow, "generate_vietnam_workbook", fail_generate)
    with pytest.raises(RuntimeError, match="Office actual failure"):
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert not list(output_root.rglob("*.xlsx"))


def test_generator_without_final_file_is_not_success(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())
    monkeypatch.setattr(workflow, "generate_vietnam_workbook", lambda _map, output, **_kw: output)

    with pytest.raises(workflow.VietnamWorkflowError, match="没有生成最终 XLSX") as error:
        workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert error.value.code == "output_missing"


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_real_generator_uses_product_creation_time_and_current_transit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())

    result = workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    book = load_workbook(result.output_xlsx, read_only=True, data_only=True)
    try:
        assert book.sheetnames == [
            "越南备货清单", "雅仓库存", "雅仓动销", "数据更改", "库存商品信息"
        ]
        main = book["越南备货清单"]
        assert main["E2"].value == "VN-A"
        assert main["AA2"].value == "2026-09-23 10:00"
        assert main["AN2"].value == 3
        assert [main[f"{column}2"].value for column in ("AO", "AP", "AQ", "AR", "AS", "AT", "AU")] == [None] * 7
        assert book["数据更改"]["A2"].value == 0.8
        assert book["数据更改"]["D2"].value == 3900
    finally:
        book.close()


@pytest.mark.skipif(
    not (os.environ.get("LXE_OFFICE_NODE") and os.environ.get("LXE_OFFICE_CLI")),
    reason="Host Office Kit paths are not configured",
)
def test_nondefault_settings_reaches_recalculated_five_sheet_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    _current(monkeypatch, _sku_map(tmp_path / "current.xlsx"))
    monkeypatch.setattr(workflow, "dataset_dir", lambda *parts: tmp_path / "artifacts" / parts[-1])
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_paths: _sources())
    store.save_parameters(store.config_json(RecommendationConfig(Decimal("0.7"), Decimal("0.6"), Decimal("0.1"), Decimal("4000"))))

    result = workflow.generate_current_vietnam_recommendation(sales="sales.xlsx", inventory="inventory.xlsx", products="products.xlsx")
    assert result.config_source == str(store.data_directory() / "parameters.json")
    assert result.config == RecommendationConfig(
        day_adjustment_30d=Decimal("0.7"), day_adjustment_15d=Decimal("0.6"),
        day_adjustment_7d=Decimal("0.1"), exchange_rate=Decimal("4000"),
    )
    book = load_workbook(result.output_xlsx, read_only=True, data_only=True)
    try:
        assert book.sheetnames == [
            "越南备货清单", "雅仓库存", "雅仓动销", "数据更改", "库存商品信息",
        ]
        assert [book["数据更改"].cell(2, column).value for column in range(1, 5)] == [
            0.7, 0.6, 0.1, 4000,
        ]
        assert [book["越南备货清单"].cell(2, column).value for column in range(48, 52)] == [
            0.7, 0.6, 0.1, 4000,
        ]
    finally:
        book.close()
