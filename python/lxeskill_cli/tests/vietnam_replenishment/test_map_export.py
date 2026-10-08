from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
import json
import os
from pathlib import Path
import subprocess
import sys
from threading import Event

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from services.vietnam_replenishment import settings
from services.vietnam_replenishment.asset_contract import load_sku_parameters
from services.vietnam_replenishment.sku_map_store import SkuMapStoreError


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))


def map_file(path: Path, cost=10):
    book = Workbook()
    book.active.append(["SKU", "热销标记", "成本", "跨境价", "折扣价"])
    book.active.append(["00123", 1, cost, 30000, 25000])
    book.active["A2"].number_format = "@"
    book.active["C2"].font = Font(bold=True, color="123456")
    book.create_sheet("用户备注").append(["保留说明", "=1+1"])
    book.save(path)
    book.close()
    return path


def test_template_is_blank_and_fillable_without_initializing_settings(tmp_path):
    output = tmp_path / "中文 模板.xlsx"
    assert settings.export_map("template", output) == {"path": str(output)}
    assert not settings.data_directory().exists()
    book = load_workbook(output)
    try:
        assert book.sheetnames == ["SKU参数映射"]
        sheet = book.active
        assert [cell.value for cell in sheet[1]] == ["SKU", "热销标记", "成本", "跨境价", "折扣价"]
        assert all(cell.value is None for row in sheet.iter_rows(min_row=2) for cell in row)
        assert sheet.freeze_panes == "A2"
        assert "1 或 2" in sheet["B1"].comment.text
        assert all("大于 0" in sheet[f"{column}1"].comment.text for column in "CDE")
        assert sheet.column_dimensions["A"].number_format == sheet["A2"].number_format == "@"
        assert all(sheet.column_dimensions[column].width >= 16 for column in "ABCDE")
        assert not any(cell.data_type == "f" for row in sheet for cell in row)
        with pytest.raises(SkuMapStoreError, match="没有 SKU"):
            settings.upload_map(output)
        for column, value in enumerate(["00123", 1, 10, 30000, 25000], 1):
            sheet.cell(2, column, value)
        book.save(output)
    finally:
        book.close()
    state = settings.upload_map(output)
    assert load_sku_parameters(Path(state["sku_map"]["path"]))["00123"].cost == 10


def test_current_export_preserves_entire_file_and_metadata_with_broken_parameters(tmp_path):
    state = settings.upload_map(map_file(tmp_path / "source.xlsx"))
    source = Path(state["sku_map"]["path"])
    params = settings.data_directory() / "parameters.json"
    params.write_text("{broken")
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (source, params)}
    for kind in ("current", "template"):
        destination = tmp_path / f"导出 {kind}.xlsx"
        destination.write_bytes(b"previous destination")
        settings.export_map(kind, destination)
        if kind == "current":
            assert destination.read_bytes() == before[source][0]
        assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in before} == before


@pytest.mark.parametrize("problem", ["missing", "corrupt"])
def test_bad_current_map_preserves_destination_and_template_remains_available(tmp_path, problem):
    if problem == "corrupt":
        settings.upload_map(map_file(tmp_path / "source.xlsx"))
        (settings.data_directory() / "sku-map.xlsx").write_bytes(b"broken zip")
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"keep this file")
    with pytest.raises(SkuMapStoreError):
        settings.export_map("current", output)
    assert output.read_bytes() == b"keep this file"
    assert not list(tmp_path.glob(".vietnam-export-*"))
    settings.export_map("template", tmp_path / "template.xlsx")


@pytest.mark.parametrize("kind", ["current", "template"])
@pytest.mark.parametrize("alias", ["direct", "symlink", "hardlink", "parent"])
def test_internal_map_cannot_be_an_export_destination(tmp_path, kind, alias):
    settings.upload_map(map_file(tmp_path / "source.xlsx"))
    source = settings.data_directory() / "sku-map.xlsx"
    before = source.read_bytes()
    output = source
    if alias in ("symlink", "hardlink"):
        output = tmp_path / "alias.xlsx"
        try:
            if alias == "symlink":
                output.symlink_to(source)
            else:
                os.link(source, output)
        except OSError as exc:
            pytest.skip(f"File alias unavailable: {exc}")
    elif alias == "parent":
        directory = tmp_path / "alias-dir"
        try:
            directory.symlink_to(source.parent, target_is_directory=True)
        except OSError as exc:
            pytest.skip(f"Directory symlink unavailable: {exc}")
        output = directory / source.name
    with pytest.raises(ValueError, match="不能覆盖"):
        settings.export_map(kind, output)
    assert source.read_bytes() == before


@pytest.mark.parametrize("kind", ["template", "current"])
def test_failed_publish_keeps_previous_destination(tmp_path, monkeypatch, kind):
    settings.upload_map(map_file(tmp_path / "source.xlsx"))
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"keep")
    def fail(*args):
        raise PermissionError("fixture: destination denied")
    monkeypatch.setattr(settings.os, "replace", fail)
    with pytest.raises(PermissionError, match="destination denied"):
        settings.export_map(kind, output)
    assert output.read_bytes() == b"keep"
    assert not list(tmp_path.glob(".vietnam-export-*"))


def test_export_takes_one_snapshot_while_upload_waits(tmp_path, monkeypatch):
    settings.upload_map(map_file(tmp_path / "original.xlsx", cost=10))
    source = settings.data_directory() / "sku-map.xlsx"
    original = source.read_bytes()
    replacement = map_file(tmp_path / "replacement.xlsx", cost=20)
    entered, release, upload_started = Event(), Event(), Event()
    copy = settings._copy_validated
    def hold_snapshot(src, dst):
        if src == source:
            entered.set()
            assert release.wait(5)
        copy(src, dst)
    def upload():
        upload_started.set()
        return settings.upload_map(replacement)
    monkeypatch.setattr(settings, "_copy_validated", hold_snapshot)
    output = tmp_path / "snapshot.xlsx"
    with ThreadPoolExecutor(max_workers=2) as pool:
        export = pool.submit(settings.export_map, "current", output)
        try:
            assert entered.wait(5)
            update = pool.submit(upload)
            assert upload_started.wait(5)
            assert not update.done()
            assert source.read_bytes() == original
        finally:
            release.set()
        export.result(timeout=5)
        update.result(timeout=5)
    assert output.read_bytes() == original
    assert source.read_bytes() == replacement.read_bytes()


def test_private_bridge_handles_unicode_paths_and_rejects_invalid_inputs(tmp_path):
    output = tmp_path / "中文 空格.xlsx"
    args = [sys.executable, "-I", "-B", "-m", "services.vietnam_replenishment.settings", "export"]
    result = subprocess.run([*args, "template", str(output)], capture_output=True, text=True, encoding="utf-8")
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["data"] == {"path": str(output)}
    for kind, path in [("invalid", str(output)), ("template", "relative.xlsx"), ("template", str(tmp_path / "other.txt"))]:
        result = subprocess.run([*args, kind, path], capture_output=True, text=True, encoding="utf-8")
        assert result.returncode == 1
        assert json.loads(result.stdout)["success"] is False


def test_partial_template_write_keeps_destination(tmp_path, monkeypatch):
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"keep")
    def fail(destination):
        destination.write_bytes(b"partial")
        raise OSError("fixture: disk full")
    monkeypatch.setattr(settings, "_write_map_template", fail)
    with pytest.raises(OSError, match="disk full"):
        settings.export_map("template", output)
    assert output.read_bytes() == b"keep"
    assert not list(tmp_path.glob(".vietnam-export-*"))
