from __future__ import annotations

import asyncio
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from services.shangman.goods_export import (
    GoodsExportWorkbookError,
    ShangmanClient,
    ShangmanCredentials,
    ShangmanAuthError,
    ShangmanDownloadUrlError,
    StaticCaptchaCodeProvider,
)


BUSINESS_HEADERS = [
    "SKU",
    "商品名",
    "总数量",
    "有效库存",
    "锁定库存",
    "在途库存",
    "预警库存",
    "7天销量",
    "15天销量",
    "30天销量",
    "仓库名称",
    "创建时间",
]

BILINGUAL_BUSINESS_HEADERS = [
    "SKU*\n(商品编码)",
    "Goods Name*\n(商品名称)",
    "Total Quantity\n(总数量)",
    "Effective Stock\n(有效库存)",
    "Lock Stock\n(锁定库存)",
    "Transit Stock\n(在途库存)",
    "Warning Quantity\n(预警库存)",
    "7 Days Sales\n(7天销量)",
    "15 Days Sales\n(15天销量)",
    "30 Days Sales\n(30天销量)",
    "Warehouse Name\n(仓库名称)",
    "Create Time\n(创建时间)",
]


class FakeResponse:
    def __init__(
        self,
        *,
        status: int = 200,
        payload: dict | None = None,
        body: bytes = b"",
        text_body: str | None = None,
    ) -> None:
        self.status = status
        self._payload = payload
        self._body = body
        self._text_body = text_body

    async def json(self, content_type=None):
        if self._payload is None:
            raise ValueError("response is not JSON")
        return self._payload

    async def text(self) -> str:
        if self._text_body is not None:
            return self._text_body
        return json.dumps(self._payload or {}, ensure_ascii=False)

    async def read(self) -> bytes:
        return self._body


class FakeRequest:
    def __init__(self, response: FakeResponse) -> None:
        self.response = response

    async def __aenter__(self) -> FakeResponse:
        return self.response

    async def __aexit__(self, exc_type, exc, traceback) -> None:
        return None


class FakeSession:
    def __init__(self, responses: list[FakeResponse]) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []

    def get(self, url: str, **kwargs) -> FakeRequest:
        self.calls.append({"method": "GET", "url": url, **kwargs})
        return FakeRequest(self.responses.pop(0))

    def post(self, url: str, **kwargs) -> FakeRequest:
        self.calls.append({"method": "POST", "url": url, **kwargs})
        return FakeRequest(self.responses.pop(0))


def xlsx_bytes(headers: list[str] = BUSINESS_HEADERS) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "sheet1"
    worksheet.append(headers)
    worksheet.append(["SKU-1", "商品", 3, 2, 1, 0, 1, 4, 8, 12, "雅加达", "2026-09-17"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def corrupt_dimension(workbook_bytes: bytes) -> bytes:
    source = BytesIO(workbook_bytes)
    output = BytesIO()
    with ZipFile(source) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for info in archive.infolist():
            content = archive.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = content.decode("utf-8").replace('ref="A1:L2"', 'ref="A1:A1"')
                content = text.encode("utf-8")
            rewritten.writestr(info, content)
    return output.getvalue()


def make_client(session: FakeSession, output_dir: Path) -> ShangmanClient:
    return ShangmanClient(
        credentials=ShangmanCredentials(
            tenant_id="tenant-1",
            username="processed-user",
            processed_password="processed-password",
            basic_auth="Basic ZHVtbXk6cGFzcw==",
        ),
        captcha_provider=StaticCaptchaCodeProvider("1234"),
        session=session,
        output_dir=output_dir,
        trusted_download_hosts={"erp.shangmanet.com", "download.example.test"},
    )


def test_export_goods_authenticates_downloads_and_returns_canonical_payload(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "captcha-key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "access-token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://download.example.test/goods.xlsx"}),
            FakeResponse(body=xlsx_bytes()),
        ]
    )

    result = asyncio.run(make_client(session, tmp_path).export_goods())

    assert result.to_payload() == {
        "platform": "上马印尼",
        "source": "shangman_goods_export",
        "artifact_path": str(tmp_path / result.filename),
        "filename": result.filename,
        "sheet_names": ["sheet1"],
        "row_count": 1,
        "headers": BUSINESS_HEADERS,
        "download_host": "download.example.test",
    }
    assert result.filename.startswith("上马印尼-商品-")
    assert result.filename.endswith(".xlsx")
    assert len(result.filename) == len("上马印尼-商品-YYYYMMDD-HHMMSS.xlsx")
    assert Path(result.artifact_path).is_file()
    assert load_workbook(result.artifact_path, read_only=True).sheetnames == ["sheet1"]

    captcha_call, login_call, export_call, download_call = session.calls
    assert captcha_call["url"].endswith("/api/blade-auth/oauth/captcha")
    assert login_call["params"] == {
        "tenantId": "tenant-1",
        "username": "processed-user",
        "password": "8e90dcfaa08d05a6b9a0e671448a7557",
        "grant_type": "captcha",
        "scope": "all",
        "type": "account",
    }
    assert login_call["headers"] == {
        "Captcha-Key": "captcha-key",
        "Captcha-Code": "1234",
        "Tenant-Id": "tenant-1",
        "Authorization": "Basic ZHVtbXk6cGFzcw==",
    }
    assert "auth" not in login_call
    assert export_call["headers"]["Blade-Auth"] == "bearer access-token"
    assert export_call["headers"]["Tenant-Id"] == "tenant-1"
    assert export_call["headers"]["Authorization"] == "Basic ZHVtbXk6cGFzcw=="
    assert "auth" not in download_call
    assert "headers" not in download_call
    assert download_call["allow_redirects"] is False


def test_export_reuses_process_local_login_state_across_clients(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "captcha-key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "shared-token", "expires_in": 600}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://download.example.test/first.xlsx"}),
            FakeResponse(body=xlsx_bytes()),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://download.example.test/second.xlsx"}),
            FakeResponse(body=xlsx_bytes()),
        ]
    )

    first = asyncio.run(make_client(session, tmp_path / "first").export_goods())
    second = asyncio.run(make_client(session, tmp_path / "second").export_goods())

    assert first.row_count == second.row_count == 1
    assert [call["method"] for call in session.calls] == [
        "GET", "POST", "POST", "GET", "POST", "GET",
    ]
    assert session.calls[1]["url"].endswith("/oauth/token")
    assert session.calls[2]["url"].endswith("/goods/merchant/exportNew")
    assert session.calls[4]["url"].endswith("/goods/merchant/exportNew")
    assert session.calls[4]["headers"]["Blade-Auth"] == "bearer shared-token"


def test_export_discards_rejected_token_and_reauthenticates_once(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "captcha-key-1", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "stale-token"}),
            FakeResponse(status=401, payload={"message": "expired"}),
            FakeResponse(payload={"key": "captcha-key-2", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "fresh-token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://download.example.test/goods.xlsx"}),
            FakeResponse(body=xlsx_bytes()),
        ]
    )

    result = asyncio.run(make_client(session, tmp_path).export_goods())

    assert result.row_count == 1
    login_calls = [call for call in session.calls if call["url"].endswith("/oauth/token")]
    export_calls = [call for call in session.calls if call["url"].endswith("/goods/merchant/exportNew")]
    assert len(login_calls) == 2
    assert [call["headers"]["Blade-Auth"] for call in export_calls] == [
        "bearer stale-token", "bearer fresh-token",
    ]


def test_workbook_validation_scans_rows_when_dimension_metadata_is_wrong(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://erp.shangmanet.com/file.xlsx"}),
            FakeResponse(body=corrupt_dimension(xlsx_bytes())),
        ]
    )

    result = asyncio.run(make_client(session, tmp_path).export_goods())

    assert result.row_count == 1
    assert result.headers == BUSINESS_HEADERS


def test_export_accepts_actual_bilingual_business_headers(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://erp.shangmanet.com/goods.xlsx"}),
            FakeResponse(body=xlsx_bytes(BILINGUAL_BUSINESS_HEADERS)),
        ]
    )

    result = asyncio.run(make_client(session, tmp_path).export_goods())

    assert result.headers == BILINGUAL_BUSINESS_HEADERS
    assert result.row_count == 1


def test_export_rejects_non_https_or_untrusted_download_url(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "http://evil.example/file.xlsx"}),
        ]
    )

    with pytest.raises(ShangmanDownloadUrlError, match="trusted https"):
        asyncio.run(make_client(session, tmp_path).export_goods())
    assert len(session.calls) == 3


def test_default_download_hosts_accept_exact_oss_host_only(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://oss.erp.shangmanet.com/goods.xlsx"}),
            FakeResponse(body=xlsx_bytes()),
        ]
    )
    client = ShangmanClient(
        credentials=ShangmanCredentials(
            tenant_id="tenant-1",
            username="processed-user",
            processed_password="processed-password",
            basic_auth="Basic ZHVtbXk6cGFzcw==",
        ),
        captcha_provider=StaticCaptchaCodeProvider("1234"),
        session=session,
        output_dir=tmp_path,
    )

    result = asyncio.run(client.export_goods())

    assert result.download_host == "oss.erp.shangmanet.com"
    download_call = session.calls[-1]
    assert download_call["url"] == "https://oss.erp.shangmanet.com/goods.xlsx"
    assert "auth" not in download_call
    assert "Blade-Auth" not in download_call.get("headers", {})
    assert download_call["allow_redirects"] is False
    with pytest.raises(ShangmanDownloadUrlError, match="trusted https"):
        client._validate_download_url("https://evil.erp.shangmanet.com/goods.xlsx")


def test_export_rejects_xlsx_without_business_headers(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://erp.shangmanet.com/file.xlsx"}),
            FakeResponse(body=xlsx_bytes(["unrelated", "headers"])),
        ]
    )

    with pytest.raises(GoodsExportWorkbookError, match="business headers"):
        asyncio.run(make_client(session, tmp_path).export_goods())
    assert not list(tmp_path.iterdir())


def test_export_rejects_malformed_xlsx_and_leaves_no_partial_artifact(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 200, "success": True, "data": "https://erp.shangmanet.com/file.xlsx"}),
            FakeResponse(body=b"not an xlsx"),
        ]
    )

    with pytest.raises(GoodsExportWorkbookError, match="readable XLSX"):
        asyncio.run(make_client(session, tmp_path).export_goods())
    assert not list(tmp_path.iterdir())


def test_export_rejects_failed_export_business_response(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(payload={"access_token": "token"}),
            FakeResponse(payload={"code": 500, "success": False, "message": "upstream failed"}),
        ]
    )

    with pytest.raises(Exception, match="upstream failed"):
        asyncio.run(make_client(session, tmp_path).export_goods())


def test_http_errors_keep_status_and_redact_credentials(tmp_path: Path) -> None:
    session = FakeSession(
        [
            FakeResponse(payload={"key": "key", "image": "data:image/png;base64,abc"}),
            FakeResponse(
                status=401,
                text_body='{"message":"password=processed-password token=access-token"}',
            ),
        ]
    )

    with pytest.raises(ShangmanAuthError) as error:
        asyncio.run(make_client(session, tmp_path).export_goods())
    assert "status=401" in str(error.value)
    assert "processed-password" not in str(error.value)
    assert "access-token" not in str(error.value)
