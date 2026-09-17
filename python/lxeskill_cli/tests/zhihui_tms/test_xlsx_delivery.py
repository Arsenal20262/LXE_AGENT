from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from services.zhihui_tms.errors import ZhihuiTmsSchemaError
from services.zhihui_tms.product_export import ZhihuiTmsExportPage, ZhihuiTmsExportResult
from services.zhihui_tms import xlsx_delivery
from services.zhihui_tms.xlsx_delivery import (
    ZhihuiTmsDeliveryError,
    deliver_product_exports,
)


def _xlsx_bytes(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class FakeDownloadClient:
    def __init__(self, payloads: dict[str, tuple[bytes, str]]) -> None:
        self.payloads = payloads
        self.calls: list[str] = []

    def download_bytes(self, url: str, **kwargs: Any) -> tuple[bytes, str]:
        self.calls.append(url)
        return self.payloads[url]


def _export_result(*pages: tuple[int, str]) -> ZhihuiTmsExportResult:
    return ZhihuiTmsExportResult(
        pages=tuple(
            ZhihuiTmsExportPage(
                page=page,
                product_ids=(page,),
                response={"code": "200", "pop": url},
            )
            for page, url in pages
        ),
        total_records=len(pages),
        reported_total_num=len(pages),
        request_count=len(pages),
    )


def test_delivers_page_artifacts_and_one_ordered_merged_workbook(tmp_path: Path) -> None:
    first_url = "https://tms-cos.mabangerp.com/export/page-1.xls"
    second_url = "https://tms-cos.mabangerp.com/export/page-2.xls"
    client = FakeDownloadClient(
        {
            first_url: (_xlsx_bytes([["SKU", "库存"], ["A", 1]]), "application/vnd.ms-excel"),
            second_url: (_xlsx_bytes([["SKU", "库存"], ["B", 2]]), "application/vnd.ms-excel"),
        }
    )

    result = deliver_product_exports(
        client,
        _export_result((1, first_url), (2, second_url)),
        output_dir=tmp_path,
        date_label="20260917",
    )

    assert [artifact.kind for artifact in result.artifacts] == ["page", "page", "merged"]
    assert [Path(artifact.path).name for artifact in result.page_artifacts] == [
        "智慧tms-商品-第1页-20260917.xlsx",
        "智慧tms-商品-第2页-20260917.xlsx",
    ]
    assert result.merged_artifact is not None
    assert Path(result.merged_artifact.path).name == "智慧tms-商品-合并-20260917.xlsx"
    assert result.headers == ("SKU", "库存")
    assert result.total_rows == 2
    assert [Path(artifact.path).is_file() for artifact in result.artifacts] == [True, True, True]

    workbook = load_workbook(result.merged_artifact.path, read_only=True, data_only=False)
    assert list(workbook.active.values) == [("SKU", "库存"), ("A", 1), ("B", 2)]
    assert client.calls == [first_url, second_url]


def test_header_mismatch_keeps_downloaded_pages_and_does_not_publish_merge(tmp_path: Path) -> None:
    first_url = "https://tms-cos.mabangerp.com/export/page-1.xls"
    second_url = "https://tms-cos.mabangerp.com/export/page-2.xls"
    client = FakeDownloadClient(
        {
            first_url: (_xlsx_bytes([["SKU", "库存"], ["A", 1]]), "application/vnd.ms-excel"),
            second_url: (_xlsx_bytes([["SKU", "数量"], ["B", 2]]), "application/vnd.ms-excel"),
        }
    )

    with pytest.raises(ZhihuiTmsDeliveryError, match="表头"):
        deliver_product_exports(
            client,
            _export_result((1, first_url), (2, second_url)),
            output_dir=tmp_path,
            date_label="20260917",
        )

    assert (tmp_path / "智慧tms-商品-第1页-20260917.xlsx").is_file()
    assert (tmp_path / "智慧tms-商品-第2页-20260917.xlsx").is_file()
    assert not (tmp_path / "智慧tms-商品-合并-20260917.xlsx").exists()


def test_download_failure_does_not_delete_prior_page_artifact(tmp_path: Path) -> None:
    first_url = "https://tms-cos.mabangerp.com/export/page-1.xls"
    second_url = "https://tms-cos.mabangerp.com/export/page-2.xls"

    class FailingClient(FakeDownloadClient):
        def download_bytes(self, url: str, **kwargs: Any) -> tuple[bytes, str]:
            self.calls.append(url)
            if url == second_url:
                raise ZhihuiTmsSchemaError("tms_download_invalid", "fixture download failed")
            return self.payloads[url]

    client = FailingClient({first_url: (_xlsx_bytes([["SKU"], ["A"]]), "application/vnd.ms-excel")})
    with pytest.raises(ZhihuiTmsSchemaError, match="fixture download failed"):
        deliver_product_exports(
            client,
            _export_result((1, first_url), (2, second_url)),
            output_dir=tmp_path,
            date_label="20260917",
        )

    assert (tmp_path / "智慧tms-商品-第1页-20260917.xlsx").is_file()
    assert not (tmp_path / "智慧tms-商品-合并-20260917.xlsx").exists()


def test_empty_export_result_returns_no_artifacts(tmp_path: Path) -> None:
    result = deliver_product_exports(
        FakeDownloadClient({}),
        _export_result(),
        output_dir=tmp_path,
        date_label="20260917",
    )

    assert result.artifacts == ()
    assert result.merged_artifact is None
    assert result.total_rows == 0


def test_date_label_must_be_a_real_yyyymmdd_value(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="YYYYMMDD"):
        deliver_product_exports(
            FakeDownloadClient({}),
            _export_result(),
            output_dir=tmp_path,
            date_label="2026-09-17",
        )


def test_xls_signature_uses_xlrd_parser(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakeSheet:
        nrows = 2

        @staticmethod
        def row_values(index: int) -> list[Any]:
            return [["SKU", "库存"], ["A", 1]][index]

    class FakeBook:
        nsheets = 1

        @staticmethod
        def sheet_by_index(index: int) -> FakeSheet:
            assert index == 0
            return FakeSheet()

        @staticmethod
        def release_resources() -> None:
            return None

    monkeypatch.setattr(xlsx_delivery.xlrd, "open_workbook", lambda **kwargs: FakeBook())

    rows = xlsx_delivery._read_workbook(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1fixture",
        source="fixture.xls",
    )

    assert rows.headers == ("SKU", "库存")
    assert rows.rows == (("A", 1),)
