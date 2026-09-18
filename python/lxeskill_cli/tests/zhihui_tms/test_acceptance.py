from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from lxeskill import cli as lxeskill_cli
from services.agent_cli.zhihui import export_products
from services.zhihui_tms.client import ZhihuiTmsClient
from services.zhihui_tms.errors import ZhihuiTmsSchemaError
from services.zhihui_tms.product_export import ZhihuiTmsExportPage, ZhihuiTmsExportResult
from shared.workspace import activate_external_workspace, activate_project_workspace


class _Response:
    def __init__(self, *, payload: dict[str, Any] | None = None, content: bytes = b"") -> None:
        self.status_code = 200
        self.payload = payload
        self.content = content
        self.text = json.dumps(payload, ensure_ascii=False) if payload is not None else ""
        self.headers = {"Content-Type": "application/octet-stream"} if content else {}

    def json(self) -> dict[str, Any]:
        assert self.payload is not None
        return self.payload

    def iter_content(self, *, chunk_size: int):
        yield self.content

    def close(self) -> None:
        pass


class _Session:
    def __init__(self, content: bytes) -> None:
        self.responses = [
            _Response(payload={"code": "200", "data": {"apiToken": "fixture-token"}}),
            _Response(payload={"code": "200", "datas": [{"id": 101}, {"id": 102}], "totalNum": 2}),
            _Response(payload={"code": "200", "pop": "https://tms-cos.mabangerp.com/fixture.xls"}),
            _Response(content=content),
        ]
        self.calls: list[tuple[str, str, dict[str, Any]]] = []

    def request(self, method: str, url: str, **kwargs: Any) -> _Response:
        self.calls.append((method, url, kwargs))
        return self.responses.pop(0)

    def close(self) -> None:
        pass


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["商品ID", "库存"])
    worksheet.append([101, 5])
    worksheet.append([102, 8])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_catalog_cli_completes_fake_login_export_download_and_xlsx_delivery(
    monkeypatch, tmp_path: Path, capsys,
) -> None:
    session = _Session(_workbook_bytes())
    monkeypatch.setenv("ZHIHUI_TMS_PRODUCTION_ENABLED", "1")
    monkeypatch.setenv("ZHIHUI_TMS_ACCOUNT", "fixture-account")
    monkeypatch.setenv("ZHIHUI_TMS_PASSWORD", "fixture-secret")
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "desktop-data"))
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    monkeypatch.setattr(
        export_products,
        "ZhihuiTmsClient",
        lambda: ZhihuiTmsClient(session=session, min_request_interval_seconds=0),
    )
    activate_external_workspace(tmp_path)
    try:
        exit_code = lxeskill_cli._main(["tms", "philippines", "products-export", "--action", "execute"])
        records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    finally:
        activate_project_workspace()

    assert exit_code == 0
    result = records[-1]
    progress = records[:-1]
    assert [record["stage"] for record in progress] == [
        "login_started", "authenticated", "listed", "exported",
        "delivery_started", "downloaded", "merged",
    ]
    assert all(record["type"] == "progress" and record["protocol_version"] == "1" for record in progress)
    assert all("fixture-secret" not in json.dumps(record, ensure_ascii=False) for record in progress)
    assert all("fixture-token" not in json.dumps(record, ensure_ascii=False) for record in progress)
    assert all("https://" not in json.dumps(record, ensure_ascii=False) for record in progress)
    assert result["ok"] is True
    assert result["data"]["total_records"] == 2
    assert result["data"]["http_attempt_count"] == 4
    assert len(result["files"]) == 1
    assert result["files"] == [item["path"] for item in result["data"]["artifacts"]]
    assert [item["kind"] for item in result["data"]["artifacts"]] == ["merged"]
    assert all(Path(path).is_file() and path.endswith(".xlsx") for path in result["files"])
    merged = load_workbook(result["files"][0], read_only=True)
    assert list(merged.active.values) == [("商品ID", "库存"), (101, 5), (102, 8)]
    assert [method for method, _url, _options in session.calls] == ["POST", "POST", "POST", "GET"]
    assert "token" not in session.calls[-1][2]["headers"]
    assert "fixture-secret" not in json.dumps(result, ensure_ascii=False)


def test_catalog_cli_failure_delivers_only_real_partial_merge(monkeypatch, tmp_path: Path, capsys) -> None:
    first_url = "https://tms-cos.mabangerp.com/fixture-1.xls"
    second_url = "https://tms-cos.mabangerp.com/fixture-2.xls"

    class FakeClient:
        request_attempt_count = 0

        def login(self, _account: str, _password: str) -> None:
            return None

        def download_bytes(self, url: str) -> tuple[bytes, str]:
            if url == first_url:
                return _workbook_bytes(), "application/vnd.ms-excel"
            raise ZhihuiTmsSchemaError("tms_download_invalid", "fixture second download failed")

    def fake_export(_client: FakeClient, **_kwargs: Any) -> ZhihuiTmsExportResult:
        return ZhihuiTmsExportResult(
            pages=(
                ZhihuiTmsExportPage(1, (101,), {"pop": first_url}),
                ZhihuiTmsExportPage(2, (102,), {"pop": second_url}),
            ),
            total_records=2,
            reported_total_num=2,
            request_count=4,
        )

    monkeypatch.setenv("ZHIHUI_TMS_PRODUCTION_ENABLED", "1")
    monkeypatch.setenv("ZHIHUI_TMS_ACCOUNT", "fixture-account")
    monkeypatch.setenv("ZHIHUI_TMS_PASSWORD", "fixture-secret")
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "desktop-data"))
    monkeypatch.delenv("LXESKILL_SKILL_SCOPE", raising=False)
    monkeypatch.setattr(export_products, "ZhihuiTmsClient", FakeClient)
    monkeypatch.setattr(export_products, "export_stockwarehouse_pages", fake_export)
    activate_external_workspace(tmp_path)
    try:
        exit_code = lxeskill_cli._main(["tms", "philippines", "products-export", "--action", "execute"])
        records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    finally:
        activate_project_workspace()

    assert exit_code != 0
    assert [record["stage"] for record in records[:-1]] == [
        "login_started", "authenticated", "delivery_started", "downloaded",
    ]
    result = records[-1]
    assert result["ok"] is False
    assert result["error"]["message"].find("fixture second download failed") >= 0
    assert result["data"]["code"] == "tms_download_invalid"
    assert len(result["files"]) == 1
    assert Path(result["files"][0]).is_file()
    assert "部分合并" in Path(result["files"][0]).name
    partial = load_workbook(result["files"][0], read_only=True)
    assert list(partial.active.values) == [("商品ID", "库存"), (101, 5), (102, 8)]
    assert not list(tmp_path.rglob("智慧tms-商品-第*页-*.xlsx"))
