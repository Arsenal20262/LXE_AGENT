from __future__ import annotations

import asyncio
import json
from dataclasses import replace
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import source_verification as verification
from services.mabang.amazon.fba import store_msku as download
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import store_msku_replenishment as replenishment
from services.mabang.amazon.fba import store_msku_sales_analysis as sales
from services.mabang.amazon.fba.sku_catalog import SkuCatalogSnapshot
from services.mabang.amazon.fba.store_resolver import FbaStore
from test_source_verification import binding, row, source
from test_replenishment_formula_sheet import case_row


def test_large_metadata_roundtrip_and_report_propagation(tmp_path, monkeypatch):
    common = {'父ASIN': 'PARENT-ASIN', '7天销量': 7, '14天销量': 14, '30天销量': 30, '90天销量': 90}
    records = [row(f'M{i}', sku=f'S{i % 1065}', **common) for i in range(1068)]
    records += [row(f'UNVERIFIED-{i:04d}-' + 'x' * 70, sku='', **common) for i in range(874)]
    skus = tuple(binding(sku=f'S{i}', kind=2 if i < 244 else 1) for i in range(1065))
    path = source(tmp_path / 'source' / '202610061353-shop_店铺MSKU数据.xlsx', records, skus)
    verified = verification.load_verified_source(path, store_name='shop')
    metadata = verified.metadata
    assert len(json.dumps(metadata['unverified_rows'], ensure_ascii=False)) > 92937
    assert verified.skus == skus
    assert verified.counts == {'original_row_count': 1942, 'binding_verified_row_count': 1068, 'binding_unverified_row_count': 874}
    book = load_workbook(path, read_only=True)
    try:
        rows = list(book[verification.SHEET].values)
        assert rows[1][0] == 'metadata_rows_v1'
        assert 'unverified_rows' not in json.loads(rows[1][1])
        assert sum(kind == 'unverified' for kind, _ in rows) == 874
        assert sum(kind == 'sku' for kind, _ in rows) == 1065
        assert all(len(value) <= 32767 for _, value in rows)
        assert book[verification.SHEET].sheet_state == 'hidden'
    finally:
        book.close()

    async def search(skus):
        return None

    async def warehouse(**kwargs):
        return tmp_path / 'mock-stock.xlsx'

    monkeypatch.setattr(inventory, 'search_warehouse_stock', search)
    monkeypatch.setattr(inventory, 'download_warehouse_stock_xlsx', warehouse)
    monkeypatch.setattr(inventory, 'parse_stock_inventory_xlsx', lambda path: {'S': Decimal(100)})
    analysis = sales.analyze_store_msku_sales('shop', input_dir=path.parent, output_dir=tmp_path / 'sales')
    stock = asyncio.run(inventory.export_store_msku_actual_inventory('shop', input_dir=path.parent, output_dir=tmp_path / 'inventory'))
    assert verification.require_matching_reports(analysis.report_xlsx_path, stock.shenzhen_warehouse_inventory_report_xlsx_path, store_name='shop') == metadata
    final = replenishment.write_replenishment_report([case_row()], tmp_path / 'final.xlsx', source_metadata=metadata)
    assert verification.read_report_metadata(final) == metadata


@pytest.mark.parametrize('excluded', [0, 3])
def test_row_format_preserves_empty_lists_order_and_duplicate_records(tmp_path, excluded):
    unknown = [row('unknown-B', sku=''), row('unknown-A', sku=''), row('unknown-B', sku='')][:excluded]
    path = source(tmp_path / 'source.xlsx', [row(), *unknown])
    metadata = verification.load_verified_source(path, store_name='shop').metadata
    assert [item['key'][0] for item in metadata['unverified_rows']] == [item['MSKU'] for item in unknown]
    book = load_workbook(path, read_only=True)
    try:
        assert sum(r[0] == 'unverified' for r in book[verification.SHEET].values) == excluded
    finally:
        book.close()


@pytest.mark.parametrize('version', [3, 4])
def test_legacy_single_cell_format_remains_readable_without_migration(tmp_path, version):
    path = source(tmp_path / 'legacy.xlsx', [row(), row('unknown', sku='')])
    original = verification.load_verified_source(path, store_name='shop')
    metadata = dict(original.metadata)
    if version == 3:
        metadata.update(version=3, shop_id='10', site='us')
        del metadata['store_id'], metadata['id_type']
        unsigned = {k: v for k, v in metadata.items() if k != 'snapshot_id'}
        metadata['snapshot_id'] = verification._digest([unsigned, [s.to_record() for s in original.skus]])
    book = load_workbook(path)
    del book[verification.SHEET]
    sheet = book.create_sheet(verification.SHEET)
    sheet.append(('kind', 'json'))
    sheet.append(('metadata', json.dumps(metadata, ensure_ascii=False)))
    for sku in original.skus:
        sheet.append(('sku', json.dumps(sku.to_record())))
    sheet.sheet_state = 'hidden'
    book.save(path)
    book.close()
    before = path.read_bytes()
    restored = verification.load_verified_source(path, store_name='shop')
    assert restored.metadata == metadata and restored.skus == original.skus
    assert path.read_bytes() == before


@pytest.mark.parametrize('kind', ['metadata_rows_v1', 'unverified', 'sku'])
def test_oversized_single_record_fails_before_replacing_sheet(kind):
    book = Workbook()
    old = book.create_sheet(verification.SHEET)
    old['A1'] = 'keep me'
    metadata = {'unverified_rows': []}
    skus = ()
    value = 'x' * 32768
    if kind == 'metadata_rows_v1':
        metadata['store_name'] = value
        record = {'store_name': value}
    elif kind == 'unverified':
        record = {'key': ['M', 'P', 'A', ''], 'reason': value}
        metadata['unverified_rows'] = [record]
    else:
        skus = (binding(sku=value),)
        record = skus[0].to_record()
    before = json.dumps(metadata)
    with pytest.raises(verification.SourceVerificationError) as caught:
        verification._write_info(book, metadata, skus)
    assert f'kind={kind}' in str(caught.value)
    assert f'cell=B{2 if kind == "metadata_rows_v1" else 3}' in str(caught.value)
    assert f'length={len(json.dumps(record, ensure_ascii=False))}' in str(caught.value)
    assert book[verification.SHEET]['A1'].value == 'keep me'
    assert json.dumps(metadata) == before
    book.close()


def test_exact_cell_limit_is_not_truncated():
    book = Workbook()
    empty_length = len(json.dumps({'note': ''}, ensure_ascii=False))
    metadata = {'unverified_rows': [], 'note': 'x' * (32767 - empty_length)}
    verification._write_info(book, metadata)
    assert len(book[verification.SHEET]['B2'].value) == 32767
    assert json.loads(book[verification.SHEET]['B2'].value)['note'] == metadata['note']
    book.close()


@pytest.mark.parametrize('mutation', ['embedded', 'duplicate_metadata', 'unknown', 'missing', 'tamper', 'truncated_legacy'])
def test_malformed_storage_is_rejected(tmp_path, mutation):
    path = source(tmp_path / 'source.xlsx', [row(), row('unknown', sku='')])
    book = load_workbook(path)
    sheet = book[verification.SHEET]
    if mutation == 'embedded':
        metadata = json.loads(sheet['B2'].value)
        metadata['unverified_rows'] = []
        sheet['B2'] = json.dumps(metadata)
    elif mutation == 'duplicate_metadata':
        sheet.append(('metadata_rows_v1', sheet['B2'].value))
    elif mutation == 'unknown':
        sheet.append(('unknown', '{}'))
    elif mutation == 'missing':
        sheet.delete_rows(3)
    elif mutation == 'tamper':
        entry = json.loads(sheet['B3'].value)
        entry['reason'] = 'changed'
        sheet['B3'] = json.dumps(entry)
    else:
        sheet['A2'] = 'metadata'
        sheet['B2'] = '{"unverified_rows": [{"key": ["' + 'x' * 32768
    book.save(path)
    book.close()
    with pytest.raises(verification.SourceVerificationError) as caught:
        verification.load_verified_source(path, store_name='shop')
    if mutation == 'truncated_legacy':
        assert 'Unterminated string' in str(caught.value)


@pytest.mark.parametrize('failure', ['exception', 'changed_metadata', 'changed_skus'])
def test_download_does_not_publish_when_source_readback_fails(tmp_path, monkeypatch, failure):
    async def stores():
        return [FbaStore('shop', '10', 'shopId')]

    async def pipeline(spec):
        staged = spec.download_file.keywords['output_dir']
        path = source(staged / '202610061353-shop_店铺MSKU数据.xlsx', [row()], None)
        return download.StoreMskuExcelResult('shop', '10', 'shopId', 1, str(path), False, False)

    async def snapshot(*args):
        return SkuCatalogSnapshot('shop', '10', 'shopId', (binding(),))

    actual_read = verification.load_verified_source
    error = verification.SourceVerificationError('actual readback failure')

    def read(path, **kwargs):
        assert not list(tmp_path.glob('*.xlsx'))
        if failure == 'exception':
            raise error
        verified = actual_read(path, **kwargs)
        if failure == 'changed_metadata':
            return replace(verified, metadata={**verified.metadata, 'collected_at': 'changed'})
        return replace(verified, skus=())

    monkeypatch.setattr(download, 'fetch_fba_stores', stores)
    monkeypatch.setattr(download, 'run_export_pipeline', pipeline)
    monkeypatch.setattr(download, 'fetch_sku_catalog_snapshot', snapshot)
    monkeypatch.setattr(verification, 'load_verified_source', read)
    with pytest.raises(verification.SourceVerificationError) as caught:
        asyncio.run(download.download_store_msku_excel('10', 'shopId', store_name='shop', output_dir=tmp_path))
    if failure == 'exception':
        assert caught.value is error
    else:
        assert '与写入前不一致' in str(caught.value)
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('failure', ['exception', 'changed_metadata'])
def test_report_stamp_does_not_succeed_after_readback_failure(tmp_path, monkeypatch, failure):
    path = source(tmp_path / 'source.xlsx', [row()])
    metadata = verification.read_report_metadata(path)
    error = verification.SourceVerificationError('actual report readback failure')

    def read(path):
        if failure == 'exception':
            raise error
        return {**metadata, 'collected_at': 'changed'}

    monkeypatch.setattr(verification, 'read_report_metadata', read)
    with pytest.raises(verification.SourceVerificationError) as caught:
        verification.stamp_report(path, metadata)
    if failure == 'exception':
        assert caught.value is error
    else:
        assert '与写入前不一致' in str(caught.value)
