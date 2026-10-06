from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.mabang.amazon.fba import store_msku_replenishment as repl
from test_mabang_store_msku_replenishment import _write_inventory_report, _write_sales_report


def set_parent(path, sheet_name, value):
    book = load_workbook(path)
    sheet = book[sheet_name]
    column = [cell.value for cell in sheet[1]].index("父ASIN") + 1
    sheet.cell(2, column).value = value
    book.save(path)
    book.close()


@pytest.mark.parametrize("parent", [None, "", " \t ", "未填写父ASIN", "PARENT-SEA"])
def test_existing_reports_calculate_with_all_parent_representations(tmp_path, parent):
    sales = _write_sales_report(tmp_path / "sales/202605251530-Amazon-Test_销量分析.xlsx")
    inventory = _write_inventory_report(tmp_path / "inventory/202605251530-Amazon-Test_真实库存（深圳仓库）.xlsx")
    set_parent(sales, "MSKU明细", parent)
    set_parent(inventory, "真实库存（深圳仓库）-组合sku", parent)
    original = {path: path.read_bytes() for path in (sales, inventory)}

    result = repl.calculate_store_msku_replenishment(
        "Amazon-Test", sales_analysis_dir=sales.parent, actual_inventory_dir=inventory.parent,
        output_dir=tmp_path / "output", unlinked_shipments_snapshot_dir=tmp_path / "snapshots",
    )

    assert result.row_count == 5
    assert Path(result.report_xlsx_path).is_file()
    assert all(path.read_bytes() == data for path, data in original.items())
    details = repl.load_sales_details(sales)
    row = repl.load_inventory_rows(inventory)[0]
    # Direct in-memory callers get the same identity as file-backed callers.
    assert repl._row_key(replace(row, parent_asin=parent or "")) in details


@pytest.mark.parametrize("field", ["msku", "asin", "local_sku", "parent_asin"])
def test_parent_normalization_does_not_hide_mismatched_product_identity(tmp_path, field):
    sales = _write_sales_report(tmp_path / "sales/202605251530-Amazon-Test_销量分析.xlsx")
    inventory = _write_inventory_report(tmp_path / "inventory/202605251530-Amazon-Test_真实库存（深圳仓库）.xlsx")
    set_parent(sales, "MSKU明细", None)
    set_parent(inventory, "真实库存（深圳仓库）-组合sku", None)
    row = replace(repl.load_inventory_rows(inventory)[0], **{field: "different"})
    with pytest.raises(repl.StoreMskuReplenishmentError, match="销量分析MSKU明细缺少匹配行"):
        repl.calculate_replenishment_rows([row], repl.load_sales_details(sales))


@pytest.mark.parametrize("other_parent", [None, " ", "未填写父ASIN"])
def test_duplicate_sales_rows_are_rejected_after_parent_normalization(tmp_path, other_parent):
    path = _write_sales_report(tmp_path / "sales.xlsx")
    book = load_workbook(path)
    sheet = book["MSKU明细"]
    column = [cell.value for cell in sheet[1]].index("父ASIN") + 1
    sheet.cell(2, column).value = None
    duplicate = [cell.value for cell in sheet[2]]
    duplicate[column - 1] = other_parent
    sheet.append(duplicate)
    book.save(path)
    book.close()
    with pytest.raises(repl.StoreMskuReplenishmentError, match="销量分析MSKU明细存在重复行"):
        repl.load_sales_details(path)
