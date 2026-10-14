from __future__ import annotations

from decimal import Decimal
import os
from pathlib import Path
from shutil import copyfile
import subprocess

import pytest
from openpyxl import Workbook, load_workbook

from services.mabang.amazon.fba import final_restock_workbook as final_workbook
from services.mabang.amazon.fba import source_verification
from services.mabang.amazon.fba import store_msku_actual_inventory as inventory
from services.mabang.amazon.fba import unlinked_shipments as shipments
from services.mabang.amazon.fba.final_restock_recommendation import final_restock_formulas
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot
from services.mabang.amazon.fba.store_msku_replenishment import StoreMskuReplenishmentError
from services.mabang.amazon.fba.final_restock_workbook import (
    FinalRestockWorkbookError,
    generate_final_restock_recommendation,
)


STORE = "Amazon-Test"
SOURCE_HEADERS = (
    "店铺名称", "MSKU", "父ASIN", "ASIN", "本地SKU", "本地SKU名称", "商品名称", "备注",
    "7天销量", "14天销量", "30天销量", "90天销量", "可售", "待入库", "预留",
    "计划入库", "在途", "待调仓", "调仓中", "单品重量(g)(cm)",
)


def _source_rows() -> list[dict[str, object]]:
    return [
        {
            "MSKU": "SEA-1", "父ASIN": "PARENT-SEA", "ASIN": "ASIN-SEA", "本地SKU": "SKU-SEA",
            "本地SKU名称": "Ocean item", "备注": "", "7天销量": 77, "14天销量": 154,
            "30天销量": 330, "90天销量": 990, "可售": 880, "待入库": 0, "预留": 0,
            "计划入库": 17, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": "50g 10*10*10",
        },
        {
            "MSKU": "CLEAR-1", "父ASIN": "PARENT-CLEAR", "ASIN": "ASIN-CLEAR", "本地SKU": "SKU-CLEAR",
            "本地SKU名称": "Clearance item", "备注": "清货", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 50,
        },
        {
            "MSKU": "COMBO-1", "父ASIN": "PARENT-COMBO", "ASIN": "ASIN-COMBO", "本地SKU": "SKU-COMBO",
            "本地SKU名称": "Combo item", "备注": "", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 25,
        },
        {
            "MSKU": "UNBOUND-1", "父ASIN": "PARENT-UNBOUND", "ASIN": "ASIN-UNBOUND", "本地SKU": "",
            "本地SKU名称": "Unbound item", "备注": "", "7天销量": 7, "14天销量": 14,
            "30天销量": 30, "90天销量": 90, "可售": 0, "待入库": 0, "预留": 0,
            "计划入库": 0, "在途": 0, "待调仓": 0, "调仓中": 0, "单品重量(g)(cm)": 50,
        },
    ]


def _write_source(path: Path, rows: list[dict[str, object]] | None = None) -> Path:
    rows = _source_rows() if rows is None else rows
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(SOURCE_HEADERS)
    for row in rows:
        sheet.append([row.get(header) for header in SOURCE_HEADERS])
    workbook.save(path)
    workbook.close()
    local_skus = tuple(
        LocalSkuDefinition(str(row["本地SKU"]), 1)
        for row in rows if str(row.get("本地SKU") or "").strip()
    )
    source_verification.annotate_source(
        path,
        SkuCatalogSnapshot(STORE, "401403", "shopId", local_skus),
        requested_store_name=STORE,
    )
    return path


def _write_inventory_report(path: Path, source_path: Path, rows: list[dict[str, object]] | None = None) -> Path:
    source = source_verification.load_verified_source(source_path, store_name=STORE)
    rows = _source_rows() if rows is None else rows
    inventory_rows = []
    for row in rows:
        child_skus = "SKU-CHILD * 2" if row["MSKU"] == "COMBO-1" else ""
        inventory_rows.append(
            inventory.ActualInventoryRow(
                msku=str(row["MSKU"]),
                parent_asin=str(row["父ASIN"]),
                asin=str(row["ASIN"]),
                local_sku=str(row["本地SKU"]),
                product_link="",
                actual_inventory=Decimal("12") if row["MSKU"] == "SEA-1" else None,
                child_skus=child_skus,
            )
        )
    inventory.write_actual_inventory_xlsx(inventory_rows, path)
    source_verification.stamp_report(path, source.metadata)
    return path


def _write_confirmed_empty_snapshot(directory: Path, *, snapshot_time: str = "202610131100") -> Path:
    query = shipments.StoreUnlinkedShipmentDownloadResult(
        store_name=STORE,
        store_id=401403,
        download_time="202610131000",
        status_results=[
            shipments.UnlinkedShipmentStatusResult(spec.status_name, 0)
            for spec in shipments.UNLINKED_SHIPMENT_STATUS_SPECS
        ],
    )
    return Path(
        shipments.build_store_unlinked_shipments_snapshot(
            [], store_name=STORE, query_result=query, snapshot_time=snapshot_time, output_dir=directory,
        ).snapshot_xlsx_path
    )


def _generate(tmp_path: Path, *, source_path: Path | None = None, inventory_path: Path | None = None,
              snapshot_path: Path | None = None):
    source_path = source_path or _write_source(
        tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx"
    )
    inventory_path = inventory_path or _write_inventory_report(
        tmp_path / "inventory.xlsx", source_path
    )
    snapshot_path = snapshot_path or _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    return generate_final_restock_recommendation(
        STORE,
        msku_source_xlsx_path=source_path,
        shenzhen_inventory_report_xlsx_path=inventory_path,
        unlinked_shipments_snapshot_path=snapshot_path,
        output_dir=tmp_path / "out",
    )


def test_generates_all_source_rows_without_formula_errors(tmp_path: Path) -> None:
    if not os.environ.get("LXE_OFFICE_NODE"):
        pytest.skip("LXE Office runtime is not configured; recalculation remains NOT VERIFIED")
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    inventory_path = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot_path = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    # Moving an unchanged, verified source must preserve the report association.
    moved_source = tmp_path / "moved" / source_path.name
    moved_source.parent.mkdir()
    copyfile(source_path, moved_source)

    result = _generate(
        tmp_path,
        source_path=moved_source,
        inventory_path=inventory_path,
        snapshot_path=snapshot_path,
    )

    output_path = Path(result.output_xlsx_path)
    assert output_path.is_file()
    assert output_path.name.endswith(f"{STORE}补货最新.xlsx")
    assert result.row_count == 4
    assert result.success is True
    assert result.formula_error_count == 0
    assert result.formula_error_examples == ()

    formulas = load_workbook(output_path, data_only=False)
    values = load_workbook(output_path, data_only=True)
    try:
        assert formulas.sheetnames == ["mskulist"]
        sheet = formulas["mskulist"]
        cached = values["mskulist"]
        headers = [sheet.cell(1, column).value for column in range(1, 35)]
        assert headers == [
            "店铺名称", "MSKU", "父ASIN", "ASIN", "7天销量", "14天销量", "30天销量", "90天销量",
            "销量趋势", "加权日销", "可售", "待入库", "预留", "计划入库", "在途", "运营中心转运",
            "运营中心处理中", "FBA总库存（马帮）", "可销售天数", "未发出货件总计", "建议运输方式",
            "空运补货天数", "海运补货天数", "空运计划", "空运实际", "空运总重", "海运计划",
            "海运实际", "海运总重", "单件重量", "本地SKU名称", "备注", "本地库存", "本地SKU",
        ]
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == "A1:AH5"
        assert sheet.sheet_view.zoomScale == 85
        assert sheet["J2"].number_format == "0.0000"
        assert sheet["S2"].number_format == "0.00"
        assert all(sheet.cell(2, col).font.name == "Calibri" for col in range(1, 35))
        assert all(not sheet.cell(2, col).font.bold for col in range(1, 35))

        formula_columns = tuple(final_restock_formulas(2))
        for row_number in range(2, 6):
            expected = final_restock_formulas(row_number)
            assert {column: sheet[f"{column}{row_number}"].value for column in formula_columns} == expected
        assert [sheet.cell(row, 2).value for row in range(2, 6)] == ["SEA-1", "CLEAR-1", "COMBO-1", "UNBOUND-1"]
        assert [sheet.cell(row, 1).value for row in range(2, 6)] == [STORE] * 4
        assert sheet["N2"].value == 17
        assert cached["T2"].value == 0
        assert sheet["AG3"].value is None
        assert sheet["AH3"].value == "SKU-CLEAR"
        assert sheet["AH4"].value == "SKU-CHILD * 2"
        assert cached["J2"].value == pytest.approx(11)
        assert cached["R2"].value == pytest.approx(880)
        assert cached["S2"].value == pytest.approx(80)
        assert cached["AC2"].value == pytest.approx(cached["AB2"].value * cached["AD2"].value)
        assert cached["AC3"].value is None  # Clearance quantity is formula-empty.
    finally:
        formulas.close()
        values.close()


def test_workbook_layout_mapping_and_pr1_formulas_before_office(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report_path = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot_path = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    verified = source_verification.load_verified_source(source_path, store_name=STORE)
    _, source_by_key = final_workbook._unique_source_rows(verified.records)
    inventory_by_key = final_workbook._load_shenzhen_rows(report_path, expected_keys=set(source_by_key))
    shipment_by_key = final_workbook._shipment_quantities(
        snapshot_path, store_name=STORE, source_data_time="202610131000"
    )
    rows = final_workbook._rows_for_workbook(
        verified.records,
        store_name=STORE,
        inventory_by_key=inventory_by_key,
        shipment_by_key=shipment_by_key,
    )
    workbook_path = tmp_path / "formulas.xlsx"
    final_workbook._write_workbook(workbook_path, rows)

    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert workbook.sheetnames == ["mskulist"]
        sheet = workbook["mskulist"]
        assert sheet.max_column == 34
        assert sheet.max_row == len(verified.records) + 1
        assert tuple(sheet.cell(1, column).value for column in range(1, 35)) == (
            "店铺名称", "MSKU", "父ASIN", "ASIN", "7天销量", "14天销量", "30天销量", "90天销量",
            "销量趋势", "加权日销", "可售", "待入库", "预留", "计划入库", "在途", "运营中心转运",
            "运营中心处理中", "FBA总库存（马帮）", "可销售天数", "未发出货件总计", "建议运输方式",
            "空运补货天数", "海运补货天数", "空运计划", "空运实际", "空运总重", "海运计划",
            "海运实际", "海运总重", "单件重量", "本地SKU名称", "备注", "本地库存", "本地SKU",
        )
        assert sheet["AH4"].value == "SKU-CHILD * 2"
        assert sheet["AH3"].value == "SKU-CLEAR"
        assert sheet["AG3"].value is None
        assert sheet["N2"].value == 17
        assert sheet["T2"].value == 0
        assert sheet["AD2"].value == 50
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref == f"A1:AH{len(rows) + 1}"
        assert sheet.sheet_view.zoomScale == 85
        assert sheet["J2"].number_format == "0.0000"
        assert sheet["S2"].number_format == "0.00"
        assert all(sheet.cell(2, column).font.name == "Calibri" for column in range(1, 35))
        assert all(not sheet.cell(2, column).font.bold for column in range(1, 35))
        for row_number in range(2, len(rows) + 2):
            expected = final_restock_formulas(row_number)
            actual = {
                column: sheet[f"{column}{row_number}"].value
                for column in expected
            }
            assert actual == expected
        assert [sheet.cell(row, 2).value for row in range(2, 6)] == ["SEA-1", "CLEAR-1", "COMBO-1", "UNBOUND-1"]
        assert [sheet.cell(1, column).value for column in range(1, 35)][29] == "单件重量"
        assert sheet["AD2"].value == 50  # Header changes; numeric values remain grams.
    finally:
        workbook.close()


@pytest.mark.parametrize("mutation", ["store", "fingerprint"])
def test_source_store_and_fingerprint_are_verified(tmp_path: Path, mutation: str) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    inventory_path = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot_path = _write_confirmed_empty_snapshot(tmp_path / "snapshots")
    if mutation == "fingerprint":
        workbook = load_workbook(source_path)
        workbook.active["I2"] = 78  # Change raw 7-day sales after source verification.
        workbook.save(source_path)
        workbook.close()
    diagnostic = "店铺不一致" if mutation == "store" else "发生变化"
    with pytest.raises(source_verification.SourceVerificationError, match=diagnostic):
        generate_final_restock_recommendation(
            "Wrong-Store" if mutation == "store" else STORE,
            msku_source_xlsx_path=source_path,
            shenzhen_inventory_report_xlsx_path=inventory_path,
            unlinked_shipments_snapshot_path=snapshot_path,
            output_dir=tmp_path / "out",
        )


@pytest.mark.parametrize("different_content", [True, False])
def test_inventory_report_must_have_the_same_source_batch(tmp_path: Path, different_content: bool) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    different_rows = _source_rows()
    if different_content:
        different_rows[0]["7天销量"] = 78
    other_source = _write_source(tmp_path / "other.xlsx", different_rows)
    source_metadata = source_verification.load_verified_source(source_path, store_name=STORE).metadata
    other_metadata = source_verification.load_verified_source(other_source, store_name=STORE).metadata
    assert source_metadata["snapshot_id"] != other_metadata["snapshot_id"]
    assert (source_metadata["source_fingerprint"] != other_metadata["source_fingerprint"]) is different_content
    report = _write_inventory_report(tmp_path / "inventory.xlsx", other_source)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(FinalRestockWorkbookError, match="批次"):
        _generate(
            tmp_path,
            source_path=source_path,
            inventory_path=report,
            snapshot_path=snapshot,
        )


def test_inventory_report_must_cover_every_source_msku(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path, _source_rows()[:2])
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(FinalRestockWorkbookError, match="覆盖与源表不一致"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_duplicate_source_msku_is_rejected(tmp_path: Path) -> None:
    rows = _source_rows()
    rows[1]["MSKU"] = rows[0]["MSKU"]
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows)
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path, rows)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(FinalRestockWorkbookError, match="重复 MSKU"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_missing_required_source_number_is_not_converted_to_zero(tmp_path: Path) -> None:
    rows = _source_rows()
    rows[0]["可售"] = None
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx", rows)
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path, rows)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots")

    with pytest.raises(FinalRestockWorkbookError, match="可售"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_snapshot_quantities_map_by_msku_and_orphans_are_rejected(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    source = source_verification.load_verified_source(source_path, store_name=STORE)
    _, source_by_key = final_workbook._unique_source_rows(source.records)
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    inventory_by_key = final_workbook._load_shenzhen_rows(report, expected_keys=set(source_by_key))

    rows = final_workbook._rows_for_workbook(
        source.records,
        store_name=STORE,
        inventory_by_key=inventory_by_key,
        shipment_by_key={"SEA-1": 3.5},
    )
    assert rows[0]["未发出货件总计"] == 3.5
    with pytest.raises(FinalRestockWorkbookError, match="源表没有的 MSKU"):
        final_workbook._rows_for_workbook(
            source.records,
            store_name=STORE,
            inventory_by_key=inventory_by_key,
            shipment_by_key={"ORPHAN": 1},
        )


def test_snapshot_from_a_different_day_is_rejected(tmp_path: Path) -> None:
    source_path = _write_source(tmp_path / "202610131000-Amazon-Test_店铺MSKU数据.xlsx")
    report = _write_inventory_report(tmp_path / "inventory.xlsx", source_path)
    snapshot = _write_confirmed_empty_snapshot(tmp_path / "snapshots", snapshot_time="202610121100")

    with pytest.raises(StoreMskuReplenishmentError, match="日期与备货数据日期不一致"):
        _generate(tmp_path, source_path=source_path, inventory_path=report, snapshot_path=snapshot)


def test_office_failure_preserves_real_diagnostic(monkeypatch, tmp_path: Path) -> None:
    def fail(arguments, *, timeout_seconds):
        assert arguments[0] == "recalculate"
        assert timeout_seconds == final_workbook.OFFICE_TIMEOUT_SECONDS
        return subprocess.CompletedProcess(arguments, 1, "", "LXE_OFFICE_NODE is not configured")

    monkeypatch.setattr(final_workbook, "run_kit_bounded", fail)
    with pytest.raises(FinalRestockWorkbookError, match="LXE_OFFICE_NODE is not configured"):
        final_workbook._recalculate(tmp_path / "in.xlsx", tmp_path / "out.xlsx")
