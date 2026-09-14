from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from services.agent_cli.yacang.export_sales_90d import run as run_sales_90d
from services.agent_cli.yacang.export_sales_monthly import run as run_sales_monthly
from services.yacang.exports.sales_90d import export_sales_90d
from services.yacang.exports.sales_monthly import export_sales_monthly
from services.yacang.projection import SALES_90D_HEADERS, SALES_MONTHLY_HEADERS
from services.yacang.validation import INVENTORY_SALES_HEADERS
from services.yacang.warehouses import WAREHOUSES


def _write_xlsx(path: Path, warehouse: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(INVENTORY_SALES_HEADERS)
    sheet.append(("SKU-1", "商品", warehouse, 1, 7, 15, 30, 60, 90, 10, 0, 0, 0, 10, 0, "2026-09-01"))
    workbook.save(path)
    workbook.close()


def _headers(path: str) -> tuple[str, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        return tuple(str(value or "") for value in next(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def _first_data_row(path: str) -> tuple[Any, ...]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        next(rows)
        return tuple(next(rows))
    finally:
        workbook.close()


class FakeClient:
    def __init__(self) -> None:
        self.login_calls: list[tuple[str, str]] = []
        self.submissions: list[tuple[int, str, str]] = []
        self.current: tuple[int, str, str] | None = None

    def login(self, mobile: str, password: str) -> None:
        self.login_calls.append((mobile, password))

    def list_downloads(self, *, limit: int = 50) -> list[dict[str, Any]]:
        if self.current is None:
            return []
        warehouse_id, start_date, end_date = self.current
        return [{
            "id": f"{warehouse_id}-{start_date}-{end_date}",
            "name": "库存动销导出",
            "type": "10",
            "path": f"https://oss-accelerate.seaya.cn/example/{warehouse_id}.xlsx",
            "param_where": json.dumps([
                ["warehouse_id", "in", [str(warehouse_id)]],
                ["create_time", ">=", 1788624000],
                ["create_time", "<", 1789315200],
            ]),
        }]

    def create_inventory_sales_export(self, *, warehouse_id: int, start_date: str, end_date: str) -> str:
        self.current = (warehouse_id, start_date, end_date)
        self.submissions.append(self.current)
        return "fake-request-id"

    def download_xlsx(self, url: str, destination: Path) -> None:
        warehouse = next(item.code for item in WAREHOUSES if f"/{item.warehouse_id}.xlsx" in url)
        _write_xlsx(destination, warehouse)


def test_exports_one_7_15_30_workbook_per_warehouse_from_four_source_requests(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_sales_monthly(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )

    assert client.login_calls == [("account", "password")]
    assert result["sales_window_days"] == [7, 15, 30]
    assert result["export_count"] == 4
    assert client.submissions == [
        (warehouse.warehouse_id, "2026-09-06", "2026-09-13")
        for warehouse in WAREHOUSES
    ]
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存动销_MY8801_2026-09-13.xlsx"
    assert Path(result["xlsx_paths"][-1]).name == "雅仓系统-库存动销_VN8806_2026-09-13.xlsx"
    assert all(_headers(path) == SALES_MONTHLY_HEADERS for path in result["xlsx_paths"])
    assert _first_data_row(result["xlsx_paths"][0]) == ("SKU-1", "商品", "MY8801", 7, 15, 30)


def test_exports_90_day_column_as_four_separate_workbooks_with_dynamic_date(tmp_path: Path) -> None:
    client = FakeClient()
    result = export_sales_90d(
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=0,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
        today=lambda: date(2026, 9, 13),
    )

    assert client.login_calls == [("account", "password")]
    assert client.submissions == [
        (warehouse.warehouse_id, "2026-09-06", "2026-09-13")
        for warehouse in WAREHOUSES
    ]
    assert result["as_of_date"] == "2026-09-13"
    assert result["sales_window_days"] == 90
    assert result["export_count"] == 4
    assert Path(result["xlsx_paths"][0]).name == "雅仓系统-库存动销-MY8801_日度90天_2026-09-13.xlsx"
    assert Path(result["xlsx_paths"][-1]).name == "雅仓系统-库存动销-VN8806_日度90天_2026-09-13.xlsx"
    assert all(_headers(path) == SALES_90D_HEADERS for path in result["xlsx_paths"])
    assert _first_data_row(result["xlsx_paths"][0]) == ("SKU-1", "商品", "MY8801", 90)


def test_second_sales_projection_reuses_the_same_four_source_workbooks(tmp_path: Path) -> None:
    client = FakeClient()
    export_sales_monthly(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=300,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )
    result = export_sales_90d(
        as_of_date="2026-09-13",
        mobile="account",
        password="password",
        output_dir=tmp_path,
        cache_max_age_seconds=300,
        client=client,  # type: ignore[arg-type]
        sleep=lambda _seconds: None,
    )

    assert len(client.login_calls) == 1
    assert len(client.submissions) == 4
    assert result["sales_window_days"] == 90
    assert {item["source"] for item in result["exports"]} == {"cache"}


@pytest.mark.parametrize(
    ("runner", "business_type"),
    [(run_sales_monthly, "sales-monthly"), (run_sales_90d, "sales-90d")],
)
def test_cli_failure_preserves_high_level_context_without_credentials(
    monkeypatch: pytest.MonkeyPatch,
    runner: Any,
    business_type: str,
) -> None:
    monkeypatch.delenv("LXE_YACANG_MOBILE", raising=False)
    monkeypatch.delenv("LXE_YACANG_PASSWORD", raising=False)
    result = runner({"as_of_date": "2026-09-13"})
    assert result["success"] is False
    assert result["business_type"] == business_type
    assert "缺少雅仓账号或密码" in result["exception"]
    assert "password" not in result["exception"]
