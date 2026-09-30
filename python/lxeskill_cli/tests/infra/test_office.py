"""Office checker and transparent launcher contracts; native QA lives in scripts/."""
import json
import os
from pathlib import Path
import signal
import subprocess
import sys
import time
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import pytest


def invoke(*arguments, cwd=None, env=None):
    return subprocess.run([sys.executable, "-m", "shared.office", *map(str, arguments)],
                          cwd=cwd, env=env, capture_output=True, text=True, timeout=15)


@pytest.mark.parametrize("extension", ["xlsx", "docx", "pptx"])
def test_check_real_document_and_assertions(tmp_path, extension):
    path = tmp_path / f"中文 input.{extension}"
    if extension == "xlsx":
        document = Workbook()
        document.active["A1"] = "中文内容"
        document.active["B1"] = 27
        document.active["B2"] = "=SUM(B1)"
    elif extension == "docx":
        document = Document()
        document.add_paragraph("中文内容")
        document.sections[0].header.paragraphs[0].text = "页眉"
    else:
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0]).shapes.title.text = "中文内容"
    document.save(path)
    source = path.read_bytes()
    report = tmp_path / "检查 report.json"
    count = ["--count", "1"] if extension != "docx" else []
    result = invoke("check", path.name, "--contains", "中文内容", "--out", report.name, *count, cwd=tmp_path)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout) == json.loads(report.read_text())
    assert json.loads(result.stdout)["verdict"] == "pass"
    missing = invoke("check", path, "--contains", "missing text")
    assert missing.returncode == 1 and json.loads(missing.stdout)["checks"][-1]["status"] == "fail"
    if extension == "xlsx":
        # Numerical cells and evaluated formulas are deliberately not asserted by --contains.
        assert invoke("check", path, "--contains", "27").returncode == 1
        assert json.loads(result.stdout)["summary"]["formulas_evaluated"] is False
    elif extension == "docx":
        assert invoke("check", path, "--count", "1").returncode == 2
    assert invoke("check", path, "--out", path).returncode == 2
    assert path.read_bytes() == source


def test_check_reports_actual_corruption_and_missing_relationship(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a ZIP")
    result = invoke("check", path)
    assert result.returncode == 1
    assert "zip" in json.loads(result.stdout)["checks"][0]["detail"].lower()
    original = tmp_path / "original.docx"
    Document().save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            if info.filename != "word/styles.xml":
                target.writestr(info, source.read(info.filename))
    result = invoke("check", path)
    assert result.returncode == 1
    assert "word/styles.xml" in json.loads(result.stdout)["checks"][0]["detail"]


def test_kit_forwards_cwd_arguments_streams_and_exit_code(tmp_path):
    cli = tmp_path / "fake cli.py"
    cli.write_text("import json,os,sys\nprint(json.dumps({'args':sys.argv[1:],'cwd':os.getcwd()}))\nprint('actual native diagnostic',file=sys.stderr)\nsys.exit(7)\n")
    env = dict(os.environ, LXE_OFFICE_NODE=sys.executable, LXE_OFFICE_CLI=str(cli))
    result = invoke("convert", "--input", "中文 file.docx", "--timeout-ms", "100", cwd=tmp_path, env=env)
    assert result.returncode == 7
    assert result.stderr == "actual native diagnostic\n"
    assert json.loads(result.stdout) == {"args": ["convert", "--input", "中文 file.docx", "--timeout-ms", "100"], "cwd": str(tmp_path)}


def test_kit_requires_explicit_existing_absolute_paths(tmp_path):
    env = dict(os.environ, LXE_OFFICE_NODE="", LXE_OFFICE_CLI="")
    result = invoke("capabilities", env=env)
    assert result.returncode == 1 and "LXE_OFFICE_NODE is not configured" in result.stderr
    env["LXE_OFFICE_NODE"] = "node"
    assert "absolute path" in invoke("capabilities", env=env).stderr
    env.update(LXE_OFFICE_NODE=sys.executable, LXE_OFFICE_CLI=str(tmp_path / "missing.js"))
    result = invoke("render", "--input", "source.xlsx", env=env)
    assert result.returncode == 1 and "missing.js" in result.stderr


@pytest.mark.skipif(os.name == "nt", reason="Windows exec cancels via taskkill /T; native smoke exercises that path")
def test_kit_forwards_cancellation_and_waits_for_child_cleanup(tmp_path):
    cli = tmp_path / "wait.py"
    ready = tmp_path / "ready"
    cli.write_text("import signal,sys,time\nfrom pathlib import Path\n"
                   "def stop(signum, frame):\n print('cancelled: actual child error',file=sys.stderr);sys.exit(23)\n"
                   "signal.signal(signal.SIGTERM,stop)\nPath('ready').write_text('ready')\nwhile True: time.sleep(.01)\n")
    process = subprocess.Popen([sys.executable, "-m", "shared.office", "render"], cwd=tmp_path,
                               env=dict(os.environ, LXE_OFFICE_NODE=sys.executable, LXE_OFFICE_CLI=str(cli)),
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    try:
        deadline = time.monotonic() + 5
        while not ready.exists() and time.monotonic() < deadline:
            time.sleep(.01)
        assert ready.exists()
        process.send_signal(signal.SIGTERM)
        stdout, stderr = process.communicate(timeout=5)
        assert process.returncode == 23 and stdout == "" and "actual child error" in stderr
    finally:
        if process.poll() is None:
            process.kill()
            process.wait()
