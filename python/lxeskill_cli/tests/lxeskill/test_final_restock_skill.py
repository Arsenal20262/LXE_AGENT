from __future__ import annotations

import json
from pathlib import Path
from unittest.mock import Mock

import pytest
from openpyxl import Workbook

from lxeskill import cli as lxeskill
from lxeskill.business import load_catalog
from services.agent_cli.mabang import generate_final_restock_recommendation as adapter
from services.mabang.amazon.fba.final_restock_workbook import FinalRestockWorkbookResult, SOURCE_REQUIRED_COLUMNS
from services.mabang.amazon.fba.sku_catalog import LocalSkuDefinition, SkuCatalogSnapshot
from services.mabang.amazon.fba.source_verification import annotate_source


COMMAND = ["replenish", "final-restock", "generate"]
SKILL = "replenishment-final-restock"
INPUTS = {
    "store_name": "Amazon-Test",
    "msku_source_xlsx_path": "source.xlsx",
    "shenzhen_inventory_report_xlsx_path": "inventory.xlsx",
    "unlinked_shipments_snapshot_path": "snapshot.xlsx",
}


def _args(inputs: dict[str, str]) -> list[str]:
    return [*COMMAND, *(item for key, value in inputs.items() for item in (f"--{key.replace('_', '-')}", value))]


def _terminal(capsys) -> dict:
    records = [json.loads(line) for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(records) == 1
    assert records[0]["type"] == "result"
    return records[0]


@pytest.fixture(autouse=True)
def isolated_scope(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", SKILL)
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))


def test_final_restock_is_exposed_with_one_owner_and_the_existing_inputs(capsys) -> None:
    entry = load_catalog()["mabang_generate_final_restock_recommendation"]
    assert entry["visibility"] == "business"
    assert entry["exposed"] is True
    assert entry["owner_skills"] == [SKILL]
    assert entry["input_modes"] == ["flags"]
    assert entry["input_schema"] == {
        "type": "object",
        "properties": {key: {"type": "string", "minLength": 1} for key in INPUTS},
        "required": list(INPUTS),
        "additionalProperties": False,
    }
    assert entry["deliver_artifacts_on_failure"] is True
    assert lxeskill.main(["list"]) == 0
    commands = _terminal(capsys)["data"]["commands"]
    assert any(item["command"] == " ".join(COMMAND) for item in commands)
    assert all(item["command"] != "replenish calculate" for item in commands)


@pytest.mark.parametrize("scope,command", [
    ("replenishment-calculate", COMMAND),
    (SKILL, ["replenish", "calculate"]),
    ("", COMMAND),
])
def test_new_and_old_commands_respect_their_own_skill_scope(scope, command, monkeypatch, capsys) -> None:
    monkeypatch.setenv("LXESKILL_SKILL_SCOPE", scope)
    assert lxeskill.main(command) == lxeskill.EXIT_ENVIRONMENT
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "skill_not_in_scope"
    assert terminal["files"] == []


@pytest.mark.parametrize("missing,diagnostic", [
    ("store_name", "store_name 不能为空"),
    ("msku_source_xlsx_path", "MSKU 源文件不存在"),
    ("shenzhen_inventory_report_xlsx_path", "深圳库存报告不存在"),
    ("unlinked_shipments_snapshot_path", "未关联货件快照不存在"),
])
def test_incomplete_inputs_return_actual_service_diagnostic(missing, diagnostic, tmp_path: Path, capsys) -> None:
    source = tmp_path / "existing-input.xlsx"
    source.touch()
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    inputs["store_name"] = "Amazon-Test"
    assert lxeskill.main(_args({key: value for key, value in inputs.items() if key != missing})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert diagnostic in terminal["error"]["message"]
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


def test_removed_source_path_flag_is_rejected_in_the_new_skill_scope(capsys) -> None:
    assert lxeskill.main([
        *_args(INPUTS), "--shenzhen-source-msku-xlsx-path", "source.xlsx",
    ]) == lxeskill.EXIT_USAGE
    terminal = _terminal(capsys)
    assert terminal["error"]["code"] == "invalid_arguments"
    assert terminal["error"]["message"] == "unknown option: --shenzhen-source-msku-xlsx-path"
    assert terminal["files"] == []


def test_source_without_native_local_inventory_is_rejected_by_cli(tmp_path: Path, capsys) -> None:
    source = tmp_path / "202610151000-Amazon-Test_店铺MSKU数据.xlsx"
    workbook = Workbook()
    columns = [column for column in SOURCE_REQUIRED_COLUMNS if column != "本地库存"]
    workbook.active.append(columns)
    workbook.active.append([
        "MSKU-1" if column == "MSKU" else "SKU-1" if column == "本地SKU" else 0
        for column in columns
    ])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    assert lxeskill.main(_args({"store_name": "Amazon-Test", **inputs})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert "本地库存" in terminal["error"]["message"]
    assert "缺少列" in terminal["error"]["message"]
    assert terminal["files"] == []


def test_skill_describes_binding_verification_advice_guard() -> None:
    skill = (Path(__file__).resolve().parents[4] / "skills" / SKILL / "SKILL.md").read_text()
    assert "`uq122`" in skill
    assert "I/J/R/S" in skill
    assert "U–AC" in skill
    assert "13 列公式" in skill
    assert "绑定核验未通过" in skill
    assert "未能匹配库存/组合目录" in skill


def test_unknown_store_returns_actual_local_source_mismatch(tmp_path: Path, capsys) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.append(["MSKU", "本地SKU"])
    workbook.active.append(["MSKU-1", "SKU-1"])
    workbook.save(source)
    workbook.close()
    annotate_source(
        source,
        SkuCatalogSnapshot("Amazon-Test", "100", "shopId", (LocalSkuDefinition("SKU-1", 1),)),
        requested_store_name="Amazon-Test",
    )
    # Store verification must fail before reading the inventory or shipment inputs.
    inputs = {key: str(source) for key in INPUTS if key != "store_name"}
    assert lxeskill.main(_args({"store_name": "Unknown-Store", **inputs})) == lxeskill.EXIT_BUSINESS
    terminal = _terminal(capsys)
    assert terminal["ok"] is False
    assert terminal["error"]["code"] == "business_cli_failed"
    assert terminal["error"]["message"] == "源数据核验店铺不一致: 请求=Unknown-Store, 快照=Amazon-Test"
    assert terminal["data"]["exception"] == terminal["error"]["message"]
    assert terminal["files"] == []


@pytest.mark.parametrize("formula_errors", [0, 1])
def test_terminal_preserves_success_or_formula_failure_and_file(
    formula_errors, tmp_path: Path, monkeypatch, capsys,
) -> None:
    output = tmp_path / "workspace" / ".lxeagent" / "artifacts" / "replenish" / "final_restock_recommendation" / "test.xlsx"
    output.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.save(output)
    workbook.close()
    output = output.resolve()
    result = FinalRestockWorkbookResult(
        store_name="Amazon-Test", source_data_time="202610131000", output_xlsx_path=str(output),
        row_count=1, formula_error_count=formula_errors,
        formula_error_examples=("AC2=#VALUE!",) if formula_errors else (),
    )
    generate = Mock(return_value=result)
    monkeypatch.setattr(adapter, "generate_final_restock_recommendation", generate)
    assert lxeskill.main(_args(INPUTS)) == (lxeskill.EXIT_BUSINESS if formula_errors else 0)
    terminal = _terminal(capsys)
    assert terminal["ok"] is (formula_errors == 0)
    assert terminal["data"] == result.to_payload()
    assert terminal["files"] == [str(output)]
    if formula_errors:
        assert terminal["error"]["message"] == result.to_payload()["exception"]
        assert "AC2=#VALUE!" in terminal["error"]["message"]
    generate.assert_called_once_with("Amazon-Test", **{key: value for key, value in INPUTS.items() if key != "store_name"})
