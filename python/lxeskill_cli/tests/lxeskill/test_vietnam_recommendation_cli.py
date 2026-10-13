from __future__ import annotations

import io
import json
from decimal import Decimal
from importlib.resources import files
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from lxeskill import cli as lxeskill
from lxeskill.business import load_catalog
from services.agent_cli.vietnam_replenishment import generate
from services.vietnam_replenishment import settings
from services.vietnam_replenishment.full_workbook import UNCHANGED
from services.vietnam_replenishment.workbook import RecommendationConfig
from shared.input_assets import current_asset, previous_asset
from shared.workspace import activate_project_workspace


COMMAND = ["vietnam", "replenishment", "calculate"]
INPUTS = {"sales_file": "sales.xlsx", "inventory_file": "inventory.xlsx", "products_file": "products.xlsx"}
ARGS = [item for name, value in INPUTS.items() for item in (f"--{name.replace('_', '-')}", value)]


def _args_with_template(tmp_path: Path) -> list[str]:
    template = tmp_path / "business-template.xlsx"
    if not template.exists():
        book = Workbook()
        book.active.title = "越南备货清单"
        for name in ("改库存", "改折扣", "报官方活动", "竞价报名", "雅仓库存", "雅仓动销", "数据更改", "库存商品信息"):
            book.create_sheet(name)
        book.save(template)
        book.close()
    return [*ARGS, "--template-file", str(template)]


def _valid_template(path: Path, marker: str) -> Path:
    resource = files("services.vietnam_replenishment").joinpath("resources/skeleton.xlsx")
    with resource.open("rb") as stream:
        book = load_workbook(stream)
    try:
        book["越南备货清单"]["B2"] = 2
        book["越南备货清单"]["E2"] = "VN-A"
        book["越南备货清单"]["G2"] = 10
        for index, name in enumerate(UNCHANGED, 1):
            book.create_sheet(name, index)["A1"] = marker
        book.save(path)
    finally:
        book.close()
    return path


def _record(capsys) -> dict:
    lines = [line for line in capsys.readouterr().out.splitlines() if line.strip()]
    assert len(lines) == 1
    return json.loads(lines[0])


def _no_workflow(monkeypatch: pytest.MonkeyPatch) -> None:
    def unexpected(**_arguments) -> None:
        raise AssertionError("workflow must not run")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", unexpected)


def test_catalog_requires_three_reports_and_allows_one_run_map() -> None:
    entry = load_catalog()["vietnam_replenishment_generate"]
    assert entry["module"] == "services.agent_cli.vietnam_replenishment.generate"
    assert entry["command_path"] == COMMAND
    assert entry["session_mode"] == "none"
    assert entry["owner_skills"] == ["vietnam-replenishment"]
    assert entry["input_schema"] == {
        "type": "object", "properties": {
            **{name: {"type": "string", "minLength": 1} for name in (*INPUTS, "sku_map_file")},
            "template_file": {"type": "string", "minLength": 1, "x-lxe-file-input": {
                "accepted_extensions": [".xlsx"],
                "instruction": "请上传完整的东南亚越南备货九表 XLSX 模板，或使用已保存的当前模板。",
            }, "x-lxe-asset-slot": "vietnam_full_workbook", "x-lxe-asset-promote-on-success": False},
        },
        "required": list(INPUTS), "additionalProperties": False,
    }
    assert entry["artifact_paths"] == [{"field": "output_xlsx", "role": "deliverable"}]
    assert entry.get("deliver_artifacts_on_failure") is not True


@pytest.mark.parametrize("arguments", [
    {}, {"map_path": "/tmp/other.xlsx"}, {**INPUTS, "cost_rate": 0.7},
    *[{key: value for key, value in INPUTS.items() if key != missing} for missing in INPUTS],
    *[{**INPUTS, key: value} for key in (*INPUTS, "sku_map_file") for value in ("", " ", 2, None)],
])
def test_adapter_rejects_missing_invalid_and_unknown_arguments(arguments: dict, monkeypatch: pytest.MonkeyPatch) -> None:
    _no_workflow(monkeypatch)
    result = generate.run(arguments)
    assert result["success"] is False
    assert result["error"]["code"] == "invalid_arguments"
    assert "output_xlsx" not in result


@pytest.mark.parametrize("mode", ["stdin", "file"])
def test_json_input_rejects_unknown_fields(
    mode: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    _no_workflow(monkeypatch)
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    input_json = json.dumps({**INPUTS, 'map_path': '/tmp/other.xlsx'})
    if mode == "stdin":
        monkeypatch.setattr("sys.stdin", io.StringIO(input_json))
        args = [*COMMAND, "--stdin-json"]
    else:
        source = tmp_path / "input.json"
        source.write_text(input_json, encoding="utf-8")
        args = [*COMMAND, "--input-json", str(source)]

    assert lxeskill.main(args) == lxeskill.EXIT_USAGE
    record = _record(capsys)
    assert record["ok"] is False
    assert record["error"]["code"] == "invalid_arguments"
    assert record["files"] == []


def test_success_delivers_only_final_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    workspace = tmp_path / "workspace"
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(workspace))
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    output = workspace / ".lxeagent" / "artifacts" / "vietnam" / "recommendations" / "run" / "越南备货清单.xlsx"
    output.parent.mkdir(parents=True)
    output.write_bytes(b"synthetic workbook")
    output = output.resolve()  # Windows may restore the directory's on-disk casing.
    monkeypatch.setattr(
        generate,
        "generate_current_vietnam_recommendation",
        lambda **_arguments: SimpleNamespace(
            output_xlsx=output, sku_count=2,
            config=RecommendationConfig(
                day_adjustment_30d=Decimal("0.7"), day_adjustment_15d=Decimal("0.6"),
                day_adjustment_7d=Decimal("0.1"), exchange_rate=Decimal("4000"),
            ),
            config_source="parameters.json", sku_map_source="sku-map.xlsx",
            source_files=({"report": "inventory-sales", "path": "sales.xlsx", "sha256": "0" * 64},),
            validation={"status": "passed", "sku_count": 2},
        ),
    )

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == 0
    record = _record(capsys)
    assert record["ok"] is True
    assert record["files"] == [str(output)]
    template_source = record["data"].pop("asset_sources")["template_file"]
    assert template_source["slot"] == "vietnam_full_workbook"
    assert template_source["from"] == "upload"
    assert template_source["file_name"] == "business-template.xlsx"
    assert record["data"] == {
        "success": True,
        "status": "completed",
        "warehouse": "VN8806",
        "sku_count": 2,
        "output_xlsx": str(output),
        "config": {
            "day_adjustment_30d": "0.7", "day_adjustment_15d": "0.6",
            "day_adjustment_7d": "0.1", "exchange_rate": "4000",
            "sales_weight_7d": "0.6", "sales_weight_15d": "0.3", "sales_weight_30d": "0.1",
        },
        "config_source": "parameters.json",
        "sku_map_source": "sku-map.xlsx",
        "source_files": [{"report": "inventory-sales", "path": "sales.xlsx", "sha256": "0" * 64}],
        "validation": {"status": "passed", "sku_count": 2},
    }
    saved = current_asset("vietnam_full_workbook")
    assert saved is None
    assert lxeskill.main([*COMMAND, *ARGS]) == lxeskill.EXIT_USAGE
    missing = _record(capsys)
    assert missing["error"]["code"] == "input_required"
    assert missing["recovery"]["field"] == "template_file"


def test_chat_template_is_temporary_across_runs_and_sessions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys,
) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_WORKSPACE_ROOT", str(tmp_path / "workspace"))
    activate_project_workspace()
    template_a = _valid_template(tmp_path / "current-a.xlsx", "A")
    template_b = _valid_template(tmp_path / "one-run-b.xlsx", "B")
    settings.upload_template(template_a)
    saved_a = current_asset("vietnam_full_workbook")
    assert saved_a is not None
    saved_mtime = saved_a.path.stat().st_mtime_ns
    assert previous_asset("vietnam_full_workbook") is None
    output = tmp_path / "workspace" / ".lxeagent" / "artifacts" / "vietnam" / "result.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(b"synthetic workbook")
    selected: list[str] = []

    def succeed(**arguments):
        selected.append(arguments["template"])
        return SimpleNamespace(
            output_xlsx=output, sku_count=1, config=RecommendationConfig(),
            config_source="parameters.json", sku_map_source=arguments["template"],
            source_files=(), validation={"status": "passed", "sku_count": 1},
        )

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", succeed)
    monkeypatch.setenv("LXE_AGENT_SESSION_ID", "chat-one")
    assert lxeskill.main([*COMMAND, *ARGS, "--template-file", str(template_b)]) == 0
    _record(capsys)
    assert selected[-1] == str(template_b)
    assert current_asset("vietnam_full_workbook") == saved_a
    assert saved_a.path.read_bytes() == template_a.read_bytes()
    assert saved_a.path.stat().st_mtime_ns == saved_mtime
    assert previous_asset("vietnam_full_workbook") is None

    monkeypatch.setenv("LXE_AGENT_SESSION_ID", "chat-two")
    assert lxeskill.main([*COMMAND, *ARGS]) == 0
    _record(capsys)
    assert selected[-1] == str(current_asset("vietnam_full_workbook").path)
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_a.read_bytes()
    assert previous_asset("vietnam_full_workbook") is None

    def fail(**_arguments):
        raise RuntimeError("synthetic calculation failure")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", fail)
    assert lxeskill.main([*COMMAND, *ARGS, "--template-file", str(template_b)]) == lxeskill.EXIT_BUSINESS
    _record(capsys)
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_a.read_bytes()
    assert previous_asset("vietnam_full_workbook") is None

    settings.upload_template(template_b)
    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", succeed)
    assert lxeskill.main([*COMMAND, *ARGS]) == 0
    _record(capsys)
    assert selected[-1] == str(current_asset("vietnam_full_workbook").path)
    assert current_asset("vietnam_full_workbook").path.read_bytes() == template_b.read_bytes()


def test_business_failure_preserves_real_redacted_diagnostic_and_delivers_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    class ObservedFailure(RuntimeError):
        code = "sku_parameter_map_required"

    def fail(**_arguments) -> None:
        raise ObservedFailure("请先上传越南 SKU 参数映射表: password=topsecret")

    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    monkeypatch.setenv("LXE_YACANG_PASSWORD", "topsecret")
    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", fail)

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == lxeskill.EXIT_BUSINESS
    record = _record(capsys)
    assert record["ok"] is False
    assert record["files"] == []
    assert record["data"]["error"]["code"] == "sku_parameter_map_required"
    assert "ObservedFailure" in record["error"]["message"]
    assert "请先上传" in record["error"]["message"]
    assert "topsecret" not in record["error"]["message"]
    assert "output_xlsx" not in record["data"]


def test_missing_generated_file_fails_without_deliverable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    output = tmp_path / "state" / "artifacts" / "vietnam" / "recommendations" / "missing.xlsx"
    monkeypatch.setattr(
        generate,
        "generate_current_vietnam_recommendation",
        lambda **_arguments: SimpleNamespace(output_xlsx=output, sku_count=2),
    )

    assert lxeskill.main([*COMMAND, *_args_with_template(tmp_path)]) == lxeskill.EXIT_BUSINESS
    record = _record(capsys)
    assert record["ok"] is False
    assert record["files"] == []
    assert "FileNotFoundError" in record["error"]["message"]


def test_reports_and_explicit_map_reach_workflow(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    expected = {**INPUTS, "sku_map_file": "my map.xlsx"}
    calls = []

    def stop_after_arguments(**arguments):
        calls.append(arguments)
        raise RuntimeError("stopped after argument forwarding")

    monkeypatch.setattr(generate, "generate_current_vietnam_recommendation", stop_after_arguments)
    args = [*_args_with_template(tmp_path), "--sku-map-file", expected["sku_map_file"]]
    assert lxeskill.main([*COMMAND, *args]) == lxeskill.EXIT_BUSINESS
    assert calls == [{**{name.removesuffix("_file"): value for name, value in expected.items()}, "template": str(tmp_path / "business-template.xlsx")}]
    assert _record(capsys)["files"] == []


def test_old_no_argument_invocation_reports_missing_paths(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    _no_workflow(monkeypatch)
    template = _args_with_template(tmp_path)[-1]
    assert lxeskill.main([*COMMAND, "--template-file", template]) == lxeskill.EXIT_BUSINESS
    result = _record(capsys)
    assert result["files"] == []
    for name in INPUTS:
        assert name in result["error"]["message"]


def test_first_calculation_requires_template_upload(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "state"))
    _no_workflow(monkeypatch)
    assert lxeskill.main([*COMMAND, *ARGS]) == lxeskill.EXIT_USAGE
    result = _record(capsys)
    assert result["error"]["code"] == "input_required"
    assert result["recovery"]["field"] == "template_file"
    assert result["files"] == []


@pytest.mark.parametrize("row_count", [2, 80])
def test_cli_aggregates_map_errors_before_reading_reports(
    tmp_path, monkeypatch, capsys, row_count,
):
    from services.vietnam_replenishment import workflow, settings

    monkeypatch.setenv("LXE_DATA_ROOT", str(tmp_path / "app"))
    monkeypatch.setattr(workflow, "load_vietnam_files", lambda **_: pytest.fail("Invalid maps must stop calculation"))
    settings.read_state()
    parameters = settings.data_directory() / "parameters.json"
    source = tmp_path / "中文 输入表.xlsx"
    book = Workbook()
    book.active.append(["SKU", "热销标记", "成本", "跨境价", "折扣价"])
    for index in range(row_count):
        book.active.append([f"VN-{index}", None, 0, 2, 3])
    book.save(source)
    book.close()
    before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (source, parameters)}
    args = [*COMMAND, *_args_with_template(tmp_path), "--sku-map-file", str(source)]
    assert lxeskill.main(args) == lxeskill.EXIT_BUSINESS
    result = _record(capsys)
    assert result["ok"] is False and result["files"] == []
    assert result["data"]["error"]["code"] == "sku_parameter_map_invalid"
    message = result["error"]["message"]
    assert f"{2 * row_count} 个错误，涉及 {row_count} 行" in message
    assert "Sheet!B2" in message and "Sheet!C2" in message
    if row_count == 2:
        assert "Sheet!B3" in message and "Sheet!C3" in message
    else:
        [report] = (tmp_path / "app" / "tmp" / "sku-map-validation").glob("*.txt")
        assert str(report) in message
        detail = report.read_text(encoding="utf-8")
        assert "Sheet!C81" in detail and detail.count("Sheet!") == 160
    assert "output_xlsx" not in result["data"]
    assert {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in before} == before
