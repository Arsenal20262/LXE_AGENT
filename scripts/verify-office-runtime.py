"""Exercise real Office files using the selected Python, Node and native Kit.

Run from the repository root with uv run --frozen python. The packaged hook runs
this script with the shipped Python and a PATH containing only shipped tools.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import signal
import subprocess
import sys
import tempfile
import time

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def fixtures(root: Path) -> dict[str, Path]:
    """Edit copies and assert that unrelated content/formatting survived."""
    picture = root / "图片 source.png"
    image = Image.new("RGB", (160, 100), "white")
    ImageDraw.Draw(image).rectangle((20, 15, 140, 85), fill="#2563eb")
    image.save(picture)
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    for row in [("项目", "数量"), ("第一项", 3), ("第二项", 5), ("合计", "=SUM(B2:B3)")]:
        ws.append(row)
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    other = wb.create_sheet("保留")
    other["A1"] = "不能丢失"
    other["A1"].font = Font(italic=True)
    original = root / "原始 表格.xlsx"
    wb.save(original)
    original_hash = hashlib.sha256(original.read_bytes()).hexdigest()
    changed = load_workbook(original)
    changed["汇总"]["B3"] = 7
    xlsx = root / "修改 表格.xlsx"
    changed.save(xlsx)
    reopened = load_workbook(xlsx)
    assert reopened.sheetnames == ["汇总", "保留"]
    assert reopened["保留"]["A1"].value == "不能丢失" and reopened["保留"]["A1"].font.italic
    assert reopened["汇总"]["A1"].font.bold and reopened["汇总"]["B4"].value == "=SUM(B2:B3)"
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_hash

    document = Document()
    document.add_heading("Office 验证", 0)
    paragraph = document.add_paragraph()
    paragraph.add_run("保留加粗").bold = True
    paragraph.add_run(" 原始文字")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(1, 0).text = "中文表格"
    document.sections[0].header.paragraphs[0].text = "页眉保留"
    document.add_picture(str(picture), width=Inches(1.5))
    doc_original = root / "原始 文档.docx"
    document.save(doc_original)
    document = Document(doc_original)
    document.paragraphs[1].runs[1].text = " 修改完成"
    docx = root / "修改 文档.docx"
    document.save(docx)
    document = Document(docx)
    assert document.paragraphs[1].runs[0].bold
    assert document.tables[0].cell(1, 0).text == "中文表格"
    assert len(document.inline_shapes) == 1
    assert document.sections[0].header.paragraphs[0].text == "页眉保留"

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(1)).text_frame
    text.paragraphs[0].add_run().text = "Office 验证"
    text.paragraphs[0].runs[0].font.bold = True
    slide.shapes.add_picture(str(picture), Inches(0.5), Inches(1.5), width=Inches(2))
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["第一季度", "第二季度"]
    data.add_series("销量", (10, 20))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1), Inches(8), Inches(5), data)
    ppt_original = root / "原始 幻灯片.pptx"
    deck.save(ppt_original)
    deck = Presentation(ppt_original)
    deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text = "Office 验证 修改完成"
    data = CategoryChartData()
    data.categories = ["第一季度", "第二季度"]
    data.add_series("销量", (42, 20))
    deck.slides[1].shapes[0].chart.replace_data(data)
    pptx = root / "修改 幻灯片.pptx"
    deck.save(pptx)
    deck = Presentation(pptx)
    assert len(deck.slides) == 2 and len(deck.slides[0].shapes) == 2
    assert deck.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold
    assert tuple(deck.slides[1].shapes[0].chart.series[0].values) == (42, 20)
    return {"xlsx": xlsx, "docx": docx, "pptx": pptx}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runtime-root", type=Path, required=True)
    parser.add_argument("--output-root", type=Path)
    args = parser.parse_args()
    root = args.output_root.resolve() if args.output_root else Path(tempfile.mkdtemp(prefix="lxe-office-验收 "))
    root.mkdir(parents=True, exist_ok=True)
    runtime = args.runtime_root.resolve()
    env = dict(os.environ, LXE_OFFICE_NODE=str(runtime / "node" / ("node.exe" if os.name == "nt" else "node")),
               LXE_OFFICE_CLI=str(runtime / "office/node_modules/@deepseek-ai/libreoffice-kit/lib/cli.js"))
    command = [sys.executable, "-m", "shared.office"]
    diagnostics = []

    def run(*arguments: str, failure: bool = False, environment: dict | None = None) -> dict:
        result = subprocess.run([*command, *map(str, arguments)], cwd=root, env=environment or env,
                                capture_output=True, text=True, encoding="utf-8", timeout=150)
        diagnostics.append({"arguments": list(map(str, arguments)), "exitCode": result.returncode,
                            "stdout": result.stdout, "stderr": result.stderr})
        (root / "diagnostics.json").write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
        assert (result.returncode != 0) == failure, diagnostics[-1]
        return json.loads(result.stderr if failure else result.stdout)

    files = fixtures(root)
    hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in root.glob("*.*") if path.suffix in {".xlsx", ".docx", ".pptx"}}
    assert run("capabilities", "--json")["runtime"]["backend"] == "native"
    rendered_images = []
    failures = []
    for extension, path in files.items():
        count = ["--count", "2"] if extension in {"xlsx", "pptx"} else []
        contains = "汇总" if extension == "xlsx" else "Office 验证"
        assert run("check", path.name, "--contains", contains, *count)["verdict"] == "pass"
        pdf = root / f"导出 {extension}.pdf"
        run("convert", "--input", path.name, "--output", pdf.name)
        assert pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 1000
        selection = ["--sheet", "汇总", "--range", "A1:C6"] if extension == "xlsx" else ["--pages", "1,2" if extension == "pptx" else "1"]
        rendered = run("render", "--input", path.name, "--output-dir", f"预览 {extension}", *selection, "--dpi", "96")
        assert rendered["images"]
        for item in rendered["images"]:
            image_path = Path(item["path"])
            image = Image.open(image_path).convert("RGBA")
            white = Image.new("RGBA", image.size, "white")
            visible = Image.alpha_composite(white, image).convert("RGB")
            # A single non-background pixel is not document content. These
            # known fixtures each contain multiple lines of text or a chart.
            histogram = visible.getcolors(visible.width * visible.height)
            foreground_pixels = visible.width * visible.height - max(count for count, _ in histogram)
            if foreground_pixels < 100:
                failures.append(f"Blank or near-uniform preview ({foreground_pixels} foreground pixels): {image_path}")
            rendered_images.append(str(image_path))
        existing_preview = root / f"预览 {extension}"
        before_images = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in existing_preview.iterdir()}
        run("render", "--input", path.name, "--output-dir", existing_preview.name, *selection, failure=True)
        assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in existing_preview.iterdir()} == before_images
        before = pdf.read_bytes()
        assert run("convert", "--input", path.name, "--output", pdf.name, failure=True)["code"]
        assert pdf.read_bytes() == before
    recalculated = root / "已重算 表格.xlsx"
    run("recalculate", "--input", files["xlsx"].name, "--output", recalculated.name)
    assert load_workbook(recalculated, data_only=False)["汇总"]["B4"].value == "=SUM(B2:B3)"
    assert load_workbook(recalculated, data_only=True)["汇总"]["B4"].value == 10
    assert load_workbook(recalculated)["保留"]["A1"].value == "不能丢失"
    assert load_workbook(recalculated)["保留"]["A1"].font.italic
    run("recalculate", "--input", files["xlsx"].name, "--output", files["xlsx"].name, failure=True)

    broken = root / "损坏 文件.docx"
    broken.write_bytes(b"not a zip archive")
    # The checker returns its failure report on stdout; the Kit returns errors on stderr.
    result = subprocess.run([*command, "check", str(broken)], env=env, capture_output=True, text=True)
    assert result.returncode == 1 and json.loads(result.stdout)["verdict"] == "fail"
    rejected = root / "不应生成.pdf"
    run("convert", "--input", broken.name, "--output", rejected.name, failure=True)
    assert not rejected.exists()
    timeout_output = root / "超时.pdf"
    error = run("convert", "--input", files["docx"].name, "--output", timeout_output.name, "--timeout-ms", "1", failure=True)
    assert error["code"] == "timeout" and not timeout_output.exists(), error
    # Exercise the real missing-engine resolver without substituting a system
    # LibreOffice. Restore the owned test payload even when the assertion fails.
    engine_platform = "win32-x64" if os.name == "nt" else "darwin-" + ("arm64" if os.uname().machine == "arm64" else "x64")
    engine = runtime / "office/node_modules/@deepseek-ai" / f"libreoffice-kit-{engine_platform}"
    hidden = engine.with_name(engine.name + ".test-hidden")
    engine.rename(hidden)
    try:
        error = run("capabilities", "--json", failure=True)
        assert error["code"] == "unavailable", error
    finally:
        hidden.rename(engine)
    cancel_source = root / "取消测试.docx"
    document = Document()
    for index in range(250):
        document.add_paragraph(f"Page {index}: cancellation keeps the original source file.")
        document.add_page_break()
    document.save(cancel_source)
    cancel_hash = hashlib.sha256(cancel_source.read_bytes()).hexdigest()
    cancel_output = root / "取消预览"
    process = subprocess.Popen([*command, "render", "--input", str(cancel_source), "--output-dir", str(cancel_output), "--max-pages", "300"],
                               cwd=root, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    time.sleep(0.5)
    assert process.poll() is None, "Cancellation fixture completed before cancellation could be tested"
    if os.name == "nt":
        # Same tree termination used by LXE exec on Windows.
        killed = subprocess.run([str(Path(os.environ["SystemRoot"]) / "System32/taskkill.exe"), "/PID", str(process.pid), "/T", "/F"], capture_output=True)
        assert killed.returncode == 0, killed.stderr
    else:
        process.send_signal(signal.SIGTERM)
    stdout, stderr = process.communicate(timeout=20)
    assert process.returncode != 0 and not stdout.strip(), (process.returncode, stdout, stderr)
    assert hashlib.sha256(cancel_source.read_bytes()).hexdigest() == cancel_hash
    diagnostics.append({"operation": "cancel", "exitCode": process.returncode, "stdout": stdout, "stderr": stderr})
    (root / "diagnostics.json").write_text(json.dumps(diagnostics, ensure_ascii=False, indent=2), encoding="utf-8")
    for path, expected in hashes.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, f"Source was modified: {path}"
    report = {"status": "failed" if failures else "passed", "failures": failures, "platform": sys.platform, "python": sys.executable, "root": str(root),
              "operations": len(diagnostics), "images": rendered_images, "formulaCache": 10}
    (root / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))
    if failures:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
